# -*- coding: utf-8 -*-
"""TempoChantier : planning et incertitudes, sans calendrier ni ressources limites."""
from __future__ import annotations

import copy
import hashlib
import io
import json
import math
import re
import zipfile
from datetime import datetime
from html import escape
from pathlib import Path
from zoneinfo import ZoneInfo

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import networkx as nx
import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from scipy import stats


APP_NAME = 'TempoChantier'
VERSION = 1
TEAL = '#00645F'
DEEP = '#084D49'
CORAL = '#D7472B'
YELLOW = '#F3C851'
INK = '#203330'
MUTED = '#516761'
PALE = '#EDF4F0'
PHASES = {
    'Préparation': '#4D8F87',
    'Structure': '#D7472B',
    'Réseaux': '#8A6BAD',
    'Finitions': '#B18613',
    'Autre': '#59787B',
}
TASK_FIELDS = (
    'id',
    'nom',
    'phase',
    'min_j',
    'mode_j',
    'max_j',
    'predecesseurs',
)
ORIGINS = (
    'Saisie utilisateur',
    'Exemple fictif',
    'Exemple fictif modifié',
    'Import : origine à documenter',
)
SECTIONS = (
    ('Planning', 'calendar_view_week'),
    ('Tâches et liens', 'account_tree'),
    ('Risque de retard', 'query_stats'),
    ('Scénarios', 'compare_arrows'),
    ('Données et rapports', 'folder_open'),
    ('Comprendre', 'school'),
)
LIMITS = (
    'Modèle pédagogique : jours de travail abstraits, liens fin-début sans décalage, '
    'ressources supposées disponibles. Aucun calendrier réel, coût, contrainte '
    'de sécurité, météo commune ou validation réglementaire.'
)
MAX_TASKS = 50
MAX_FILE_BYTES = 2*1024*1024


def finite_number(value, name, lower, upper, integer=False):
    if isinstance(value, (bool, np.bool_)):
        raise ValueError(f'{name} : un nombre est requis, pas une case cochée.')
    if not isinstance(value, (int, float, np.integer, np.floating)):
        raise ValueError(f'{name} : renseigner un nombre.')
    number = float(value)
    if not math.isfinite(number) or not lower <= number <= upper:
        raise ValueError(f'{name} : valeur attendue entre {lower:g} et {upper:g}.')
    if integer and int(number) != number:
        raise ValueError(f'{name} : un entier est requis.')
    return int(number) if integer else number


def clean_text(value, name, maximum=100, allow_empty=False):
    if not isinstance(value, str):
        raise ValueError(f'{name} : renseigner un texte.')
    value = value.strip()
    if not allow_empty and not value:
        raise ValueError(f'{name} : le texte ne peut pas être vide.')
    if len(value) > maximum or any(ord(char) < 32 for char in value):
        raise ValueError(f'{name} : texte simple, limité à {maximum} caractères.')
    return value


def parse_predecessors(value):
    if value is None:
        return []
    if isinstance(value, (float, np.floating)) and math.isnan(value):
        return []
    if isinstance(value, str):
        values = re.split(r'[,;\s]+', value.strip()) if value.strip() else []
    elif isinstance(value, list):
        values = value
    else:
        raise ValueError('Prédécesseurs : utiliser des identifiants séparés par des virgules.')
    output = []
    for item in values:
        ident = clean_text(item, 'Prédécesseur', 16)
        if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,15}', ident):
            raise ValueError('Un identifiant utilise lettres sans accents, chiffres, tiret ou soulignement.')
        if ident in output:
            raise ValueError(f'Lien répété vers {ident}.')
        output.append(ident)
    if len(output) > 10:
        raise ValueError('Dix prédécesseurs au maximum par tâche dans cet atelier.')
    return output


def validate_tasks(rows):
    if not isinstance(rows, list) or len(rows) > MAX_TASKS:
        raise ValueError(f'Le planning doit contenir au maximum {MAX_TASKS} tâches.')
    clean, identifiers = [], set()
    for index, row in enumerate(rows, 1):
        if not isinstance(row, dict) or set(row) != set(TASK_FIELDS):
            raise ValueError(f'Ligne {index} : colonnes attendues {", ".join(TASK_FIELDS)}.')
        ident = clean_text(row['id'], f'Ligne {index}, ID', 16)
        if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,15}', ident):
            raise ValueError(f'{ident} : identifiant sans espace ni accent requis.')
        if ident in identifiers:
            raise ValueError(f'Identifiant dupliqué : {ident}.')
        identifiers.add(ident)
        name = clean_text(row['nom'], f'{ident}, nom', 80)
        phase = clean_text(row['phase'], f'{ident}, phase', 30)
        if phase not in PHASES:
            raise ValueError(f'{ident} : choisir une phase de la liste.')
        durations = {
            field: finite_number(row[field], f'{ident}, {field}', 0, 365)
            for field in ('min_j', 'mode_j', 'max_j')
        }
        if not durations['min_j'] <= durations['mode_j'] <= durations['max_j']:
            raise ValueError(f'{ident} : durée optimiste ≤ probable ≤ pessimiste requise.')
        preds = parse_predecessors(row['predecesseurs'])
        if ident in preds:
            raise ValueError(f'{ident} ne peut pas dépendre d’elle-même.')
        clean.append(dict(id=ident, nom=name, phase=phase, **durations, predecesseurs=preds))
    graph = nx.DiGraph()
    graph.add_nodes_from(task['id'] for task in clean)
    for task in clean:
        for pred in task['predecesseurs']:
            if pred not in identifiers:
                raise ValueError(f"{task['id']} : prédécesseur inconnu {pred}.")
            graph.add_edge(pred, task['id'])
    if not nx.is_directed_acyclic_graph(graph):
        cycle = nx.find_cycle(graph)
        names = [edge[0] for edge in cycle]+[cycle[0][0]]
        raise ValueError('Dépendance circulaire : '+' → '.join(names)+'.')
    return clean


def new_project():
    return {
        'application': APP_NAME,
        'version': VERSION,
        'titre': 'Mon planning de chantier',
        'origine': 'Saisie utilisateur',
        'taches': [],
        'parametres': {
            'echeance_j': 20.,
            'tirages': 10000,
            'graine': 42,
            'base': 'mode',
        },
        'variante': {
            'id': None,
            'min_j': 0.,
            'mode_j': 0.,
            'max_j': 0.,
        },
        'simulation_active': False,
        'comparaison_active': False,
    }


def validate_project(project):
    if not isinstance(project, dict) or set(project) != set(new_project()):
        raise ValueError('Structure de projet TempoChantier incomplète ou inconnue.')
    c = copy.deepcopy(project)
    if c['application'] != APP_NAME or type(c['version']) is not int or c['version'] != VERSION:
        raise ValueError('Fichier TempoChantier version 1 requis.')
    c['titre'] = clean_text(c['titre'], 'Nom du planning', 100)
    if c['origine'] not in ORIGINS:
        raise ValueError('Origine des données inconnue.')
    c['taches'] = validate_tasks(c['taches'])
    p = c['parametres']
    if not isinstance(p, dict) or set(p) != {'echeance_j', 'tirages', 'graine', 'base'}:
        raise ValueError('Paramètres du projet incomplets.')
    c['parametres'] = {
        'echeance_j': finite_number(p['echeance_j'], 'Échéance', 0, 20000),
        'tirages': finite_number(p['tirages'], 'Tirages', 100, 50000, True),
        'graine': finite_number(p['graine'], 'Graine', 0, 2147483647, True),
        'base': p['base'],
    }
    if p['base'] not in ('mode', 'moyenne'):
        raise ValueError('Base nominale : mode ou moyenne triangulaire attendue.')
    variant = c['variante']
    if not isinstance(variant, dict) or set(variant) != {'id', 'min_j', 'mode_j', 'max_j'}:
        raise ValueError('Variante incomplète.')
    ids = {task['id'] for task in c['taches']}
    if variant['id'] is not None and variant['id'] not in ids:
        raise ValueError('La tâche choisie pour la variante n’existe plus.')
    for field in ('min_j', 'mode_j', 'max_j'):
        variant[field] = finite_number(variant[field], 'Variante '+field, 0, 365)
    if not variant['min_j'] <= variant['mode_j'] <= variant['max_j']:
        raise ValueError('Variante : minimum ≤ probable ≤ maximum requis.')
    for field in ('simulation_active', 'comparaison_active'):
        if type(c[field]) is not bool:
            raise ValueError('États de calcul booléens requis.')
        if c[field] and not c['taches']:
            raise ValueError('Enregistrer des tâches avant une analyse.')
    if c['comparaison_active'] and variant['id'] is None:
        raise ValueError('Choisir la tâche de la variante.')
    return c


def task(ident, name, phase, low, mode, high, predecessors=''):
    return {
        'id': ident,
        'nom': name,
        'phase': phase,
        'min_j': low,
        'mode_j': mode,
        'max_j': high,
        'predecesseurs': parse_predecessors(predecessors),
    }


def example_project(key='chantier'):
    c = new_project()
    c['origine'] = 'Exemple fictif'
    if key == 'sept_jours':
        c['titre'] = 'Mini-lot : 7 jours'
        c['parametres']['echeance_j'] = 7.
        c['taches'] = [
            task('A', 'Préparation', 'Préparation', 2, 2, 2),
            task('B', 'Atelier coffrage', 'Structure', 4, 4, 4, 'A'),
            task('C', 'Atelier armatures', 'Structure', 3, 3, 3, 'A'),
            task('D', 'Assemblage', 'Finitions', 1, 1, 1, 'B,C'),
        ]
    elif key == 'deux_chemins':
        c['titre'] = 'Deux chemins concurrents'
        c['parametres']['echeance_j'] = 7.
        c['taches'] = [
            task('A', 'Installation', 'Préparation', 1, 1, 1),
            task('B', 'Atelier B', 'Structure', 2, 4, 6, 'A'),
            task('C', 'Atelier C', 'Réseaux', 2, 4, 6, 'A'),
            task('D', 'Réunion des lots', 'Finitions', 1, 1, 1, 'B,C'),
        ]
    elif key == 'chantier':
        c['titre'] = 'Petit chantier pédagogique'
        c['parametres']['echeance_j'] = 20.
        c['taches'] = [
            task('A', 'Installation du chantier', 'Préparation', 1, 2, 3),
            task('B', 'Terrassement', 'Préparation', 2, 3, 5, 'A'),
            task('C', 'Préparation du coffrage', 'Structure', 1, 2, 3, 'A'),
            task('D', 'Fondations', 'Structure', 3, 4, 6, 'B,C'),
            task('E', 'Réseaux', 'Réseaux', 1, 2, 4, 'D'),
            task('F', 'Élévation', 'Structure', 3, 5, 7, 'D'),
            task('G', 'Finitions', 'Finitions', 2, 3, 5, 'E,F'),
            task('H', 'Réception pédagogique', 'Finitions', .5, 1, 2, 'G'),
        ]
    else:
        raise ValueError('Exemple inconnu.')
    return validate_project(c)


def canonical(project):
    return json.dumps(project, ensure_ascii=False, sort_keys=True, allow_nan=False)


def make_graph(tasks):
    graph = nx.DiGraph()
    for row in tasks:
        graph.add_node(row['id'])
    for row in tasks:
        graph.add_edges_from((pred, row['id']) for pred in row['predecesseurs'])
    return graph


def numerical_tolerance(value):
    return 64*np.finfo(float).eps*np.maximum(1., np.abs(value))


def late_mask(finishes, deadline):
    threshold = deadline+8*np.spacing(abs(deadline)) if deadline else 0.
    return np.asarray(finishes, dtype=float) > threshold


def schedule_from_durations(project, durations):
    graph = make_graph(project['taches'])
    order = list(nx.lexicographical_topological_sort(graph))
    starts, ends = {}, {}
    for node in order:
        starts[node] = max((ends[pred] for pred in graph.predecessors(node)), default=0.)
        ends[node] = starts[node]+float(durations[node])
    finish = max(ends.values())
    latest_start, latest_end = {}, {}
    for node in reversed(order):
        latest_end[node] = min((latest_start[s] for s in graph.successors(node)), default=finish)
        latest_start[node] = latest_end[node]-float(durations[node])
    weighted = nx.DiGraph()
    weighted.add_nodes_from(['__start', '__finish']+order)
    for node in order:
        if graph.in_degree(node) == 0:
            weighted.add_edge('__start', node, weight=float(durations[node]))
        for parent in graph.predecessors(node):
            weighted.add_edge(parent, node, weight=float(durations[node]))
        if graph.out_degree(node) == 0:
            weighted.add_edge(node, '__finish', weight=0.)
    longest = nx.dag_longest_path_length(weighted, weight='weight')
    if not math.isclose(finish, longest, abs_tol=float(numerical_tolerance(finish))):
        raise ArithmeticError('Contrôle du chemin critique incohérent.')
    sink = min((node for node in order if graph.out_degree(node) == 0), key=lambda n: (-ends[n], n))
    path = [sink]
    while graph.in_degree(path[-1]):
        previous = min(graph.predecessors(path[-1]), key=lambda n: (-ends[n], n))
        path.append(previous)
    path.reverse()
    by_id = {row['id']: row for row in project['taches']}
    rows = []
    epsilon = float(numerical_tolerance(finish))
    for node in order:
        row = by_id[node]
        slack = max(0., latest_start[node]-starts[node])
        free = min((starts[s] for s in graph.successors(node)), default=finish)-ends[node]
        rows.append({
            **row,
            'duree_j': float(durations[node]),
            'debut_j': starts[node],
            'fin_j': ends[node],
            'debut_tard_j': latest_start[node],
            'fin_tard_j': latest_end[node],
            'marge_j': 0. if slack <= epsilon else slack,
            'marge_libre_j': 0. if abs(free) <= epsilon else max(0., free),
            'critique': slack <= epsilon,
            'successeurs': sorted(graph.successors(node)),
        })
    table = pd.DataFrame(rows)
    return {
        'table': table,
        'graph': graph,
        'order': order,
        'duree_j': finish,
        'chemin': path,
        'n_critiques': int(table['critique'].sum()),
        'somme_durees_j': float(sum(durations.values())),
        'controle_networkx': float(longest),
    }


def compute_schedule(project):
    c = validate_project(project)
    if not c['taches']:
        raise ValueError('Aucune tâche enregistrée.')
    durations = {}
    for row in c['taches']:
        value = row['mode_j'] if c['parametres']['base'] == 'mode' else (row['min_j']+row['mode_j']+row['max_j'])/3
        durations[row['id']] = value
    return schedule_from_durations(c, durations)


def binomial_interval(k, n):
    n = finite_number(n, 'Nombre de tirages', 1, 50000, True)
    k = finite_number(k, 'Dépassements', 0, n, True)
    low = 0. if k == 0 else float(stats.beta.ppf(.025, k, n-k+1))
    high = 1. if k == n else float(stats.beta.ppf(.975, k+1, n-k))
    return low, high


def sample_task_durations(tasks, n_draws, seed):
    ids = sorted(row['id'] for row in tasks)
    by_id = {row['id']: row for row in tasks}
    uniforms = np.random.default_rng(seed).random((n_draws, len(ids)))
    samples = {}
    for index, ident in enumerate(ids):
        row = by_id[ident]
        low, mode, high = row['min_j'], row['mode_j'], row['max_j']
        if high == low:
            values = np.full(n_draws, low)
        else:
            shape = (mode-low)/(high-low)
            values = stats.triang.ppf(uniforms[:, index], shape, loc=low, scale=high-low)
        samples[ident] = np.asarray(values, dtype=float)
    return samples


def simulate_project(project):
    c = validate_project(project)
    if not c['taches']:
        raise ValueError('Enregistrer les tâches avant de simuler.')
    p = c['parametres']
    n = p['tirages']
    base = compute_schedule(c)
    graph, order = base['graph'], base['order']
    samples = sample_task_durations(c['taches'], n, p['graine'])
    starts, ends = {}, {}
    zeros = np.zeros(n)
    for node in order:
        parents = list(graph.predecessors(node))
        starts[node] = np.maximum.reduce([ends[pred] for pred in parents]) if parents else zeros
        ends[node] = starts[node]+samples[node]
    finish = np.maximum.reduce([ends[node] for node in order])
    latest_start = {}
    epsilon = numerical_tolerance(finish)
    criticality = {}
    for node in reversed(order):
        successors = list(graph.successors(node))
        latest_end = np.minimum.reduce([latest_start[s] for s in successors]) if successors else finish
        latest_start[node] = latest_end-samples[node]
        criticality[node] = float(np.mean(np.abs(latest_start[node]-starts[node]) <= epsilon))
    mask = late_mask(finish, p['echeance_j'])
    count = int(mask.sum())
    low, high = binomial_interval(count, n)
    support_low = schedule_from_durations(c, {row['id']: row['min_j'] for row in c['taches']})['duree_j']
    support_high = schedule_from_durations(c, {row['id']: row['max_j'] for row in c['taches']})['duree_j']
    constant = abs(support_high-support_low) <= numerical_tolerance(support_high)
    by_id = {row['id']: row for row in c['taches']}
    influence_rows = []
    variable_finish = float(np.ptp(finish)) > float(numerical_tolerance(np.max(finish)))
    for node in order:
        duration = samples[node]
        variable_duration = float(np.ptp(duration)) > float(numerical_tolerance(np.max(duration)))
        correlation = None
        if variable_finish and variable_duration:
            correlation = float(stats.spearmanr(duration, finish).statistic)
        influence_rows.append({
            'id': node,
            'nom': by_id[node]['nom'],
            'criticite_pct': 100*criticality[node],
            'correlation_rang': correlation,
            'duree_moyenne_simulee_j': float(np.mean(duration)),
            'duree_moyenne_theorique_j': (by_id[node]['min_j']+by_id[node]['mode_j']+by_id[node]['max_j'])/3,
        })
    quantiles = np.quantile(finish, [.05, .5, .8, .9, .95])
    frames = pd.DataFrame({
        'tirage': np.arange(1, n+1),
        'duree_projet_j': finish,
        'retard_j': np.maximum(0., finish-p['echeance_j']),
        'depasse_echeance': mask,
    })
    return {
        'durees': finish,
        'table': frames,
        'influence': pd.DataFrame(influence_rows),
        'n': n,
        'depassements': count,
        'probabilite': count/n,
        'ic_bas': low,
        'ic_haut': high,
        'moyenne_j': float(np.mean(finish)),
        'ecart_type_j': float(np.std(finish, ddof=1)),
        'quantiles': dict(zip(['P5','P50','P80','P90','P95'], map(float, quantiles))),
        'support_min_j': float(support_low),
        'support_max_j': float(support_high),
        'constant': bool(constant),
        'nominal': base,
    }


def variant_project(project):
    c = validate_project(project)
    variant = c['variante']
    if variant['id'] is None:
        raise ValueError('Choisir une tâche pour la variante.')
    for row in c['taches']:
        if row['id'] == variant['id']:
            row.update({field: variant[field] for field in ('min_j','mode_j','max_j')})
    c['comparaison_active'] = False
    return validate_project(c)


def compare_projects(project):
    c = validate_project(project)
    other = variant_project(c)
    current = simulate_project(c)
    modified = simulate_project(other)
    rows = []
    indicators = [
        ('Durée nominale (j)', current['nominal']['duree_j'], modified['nominal']['duree_j']),
        ('Durée moyenne simulée (j)', current['moyenne_j'], modified['moyenne_j']),
        ('P80 (j)', current['quantiles']['P80'], modified['quantiles']['P80']),
        ('P95 (j)', current['quantiles']['P95'], modified['quantiles']['P95']),
        ('Risque de retard (%)', current['probabilite']*100, modified['probabilite']*100),
    ]
    for title, a, b in indicators:
        rows.append({'Indicateur': title, 'Actuel': a, 'Variante': b, 'Écart variante - actuel': b-a})
    differences = modified['durees']-current['durees']
    return {
        'actuel': current,
        'variante': modified,
        'table': pd.DataFrame(rows),
        'difference_moyenne_j': float(np.mean(differences)),
        'part_plus_courte': float(np.mean(differences < -numerical_tolerance(current['durees']))),
    }


def deadline_table(simulation, deadlines):
    rows = []
    for day in deadlines:
        value = finite_number(day, 'Échéance à comparer', 0, 20000)
        count = int(late_mask(simulation['durees'], value).sum())
        low, high = binomial_interval(count, simulation['n'])
        rows.append({
            'echeance_j': value,
            'risque_pct': count/simulation['n']*100,
            'ic_bas_pct': low*100 if not simulation['constant'] else None,
            'ic_haut_pct': high*100 if not simulation['constant'] else None,
        })
    return pd.DataFrame(rows)


def concurrency_table(schedule):
    frame = schedule['table']
    breaks = sorted({0., *frame['debut_j'], *frame['fin_j']})
    rows = []
    for start, end in zip(breaks[:-1], breaks[1:]):
        if end <= start:
            continue
        middle = (start+end)/2
        active = frame.loc[(frame['debut_j'] <= middle) & (frame['fin_j'] > middle), 'id'].tolist()
        rows.append({'debut_j': start, 'fin_j': end, 'taches_simultanees': len(active), 'ids': ', '.join(active)})
    return pd.DataFrame(rows, columns=['debut_j','fin_j','taches_simultanees','ids'])


@st.cache_data(max_entries=8, ttl=1800, show_spinner=False)
def cached_schedule(configuration):
    return compute_schedule(json.loads(configuration))


@st.cache_data(max_entries=6, ttl=1800, show_spinner=False)
def cached_simulation(configuration):
    return simulate_project(json.loads(configuration))


@st.cache_data(max_entries=4, ttl=1800, show_spinner=False)
def cached_comparison(configuration):
    return compare_projects(json.loads(configuration))


def tasks_frame(tasks):
    rows = []
    for task_row in tasks:
        row = copy.deepcopy(task_row)
        row['predecesseurs'] = ', '.join(row['predecesseurs'])
        rows.append(row)
    if rows:
        return pd.DataFrame(rows, columns=TASK_FIELDS)
    return pd.DataFrame({
        'id': pd.Series(dtype='str'),
        'nom': pd.Series(dtype='str'),
        'phase': pd.Series(dtype='str'),
        'min_j': pd.Series(dtype='float64'),
        'mode_j': pd.Series(dtype='float64'),
        'max_j': pd.Series(dtype='float64'),
        'predecesseurs': pd.Series(dtype='str'),
    })


def safe_csv(frame):
    clean = frame.copy()
    for column in clean:
        if pd.api.types.is_string_dtype(clean[column]) or clean[column].dtype == object:
            clean[column] = clean[column].map(
                lambda value: "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value
            )
    return clean.to_csv(index=False, float_format='%.8f').encode('utf-8-sig')


def import_tasks(content, filename):
    if len(content) > MAX_FILE_BYTES:
        raise ValueError('Fichier limité à 2 Mo.')
    if filename.lower().endswith('.csv'):
        frame = pd.read_csv(io.BytesIO(content), dtype={'id':str}, nrows=MAX_TASKS+1)
    elif filename.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20*1024*1024:
                raise ValueError('Classeur décompressé trop volumineux.')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet_name = 'Taches' if 'Taches' in workbook.sheetnames else workbook.sheetnames[0]
        workbook.close()
        frame = pd.read_excel(io.BytesIO(content), sheet_name=sheet_name, dtype={'id':str}, nrows=MAX_TASKS+1)
    else:
        raise ValueError('Utiliser un fichier CSV ou Excel .xlsx.')
    required = {'id','nom','min_j','mode_j','max_j'}
    if not required.issubset(frame.columns):
        raise ValueError('Colonnes requises : '+', '.join(sorted(required))+'.')
    if not 1 <= len(frame) <= MAX_TASKS:
        raise ValueError(f'Importer entre 1 et {MAX_TASKS} tâches.')
    if 'phase' not in frame:
        frame['phase'] = 'Autre'
    if 'predecesseurs' not in frame:
        frame['predecesseurs'] = ''
    frame['predecesseurs'] = frame['predecesseurs'].fillna('')
    return validate_tasks(frame[list(TASK_FIELDS)].astype(object).to_dict('records'))


def import_project(content):
    if len(content) > MAX_FILE_BYTES:
        raise ValueError('Fichier limité à 2 Mo.')
    def reject_constant(value):
        raise ValueError('Valeur JSON non finie interdite : '+value)
    obj = json.loads(content.decode('utf-8-sig'), parse_constant=reject_constant)
    return validate_project(obj)


def candidate_from_tasks(project, rows, origin='Saisie utilisateur'):
    c = copy.deepcopy(project)
    c.update(taches=rows, origine=origin, simulation_active=False, comparaison_active=False)
    c['variante'] = {'id':None, 'min_j':0., 'mode_j':0., 'max_j':0.}
    return validate_project(c)


def result_sheets(project):
    c = validate_project(project)
    sheets = {
        'Projet': pd.DataFrame([
            {'Paramètre':'Titre', 'Valeur':c['titre']},
            {'Paramètre':'Origine', 'Valeur':c['origine']},
            {'Paramètre':'Unité', 'Valeur':'Jour de travail abstrait'},
            *[{'Paramètre':key, 'Valeur':str(value)} for key,value in c['parametres'].items()],
        ]),
        'Taches': tasks_frame(c['taches']),
    }
    if c['taches']:
        result = cached_schedule(canonical(c))
        frame = result['table'].copy()
        for field in ('predecesseurs', 'successeurs'):
            frame[field] = frame[field].map(lambda value:', '.join(value))
        sheets['Calculs_CPM'] = frame
        sheets['Parallelisme'] = concurrency_table(result)
    if c['simulation_active']:
        sim = cached_simulation(canonical(c))
        sheets['Simulation'] = pd.DataFrame([{
            'tirages': sim['n'],
            'risque_pct': 100*sim['probabilite'],
            'moyenne_j': sim['moyenne_j'],
            **sim['quantiles'],
        }])
        sheets['Influence'] = sim['influence']
    if c['comparaison_active']:
        sheets['Comparaison'] = cached_comparison(canonical(c))['table']
    return sheets


@st.cache_data(max_entries=5, ttl=1800, show_spinner=False)
def excel_bytes(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in sheets.items():
        sheet = workbook.create_sheet(name)
        sheet.append([str(column) for column in frame.columns])
        for row in frame.itertuples(index=False, name=None):
            values = []
            for value in row:
                if isinstance(value, np.generic):
                    value = value.item()
                if isinstance(value, float) and not math.isfinite(value):
                    value = None
                values.append(value)
            sheet.append(values)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor=TEAL[1:])
            cell.font = Font(name='Calibri', color='FFFFFF', bold=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(14, len(str(column[0].value))+2))
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def fmt(value, digits=2):
    if value is None or not math.isfinite(float(value)):
        return '—'
    return f'{value:,.{digits}f}'.replace(',', ' ').replace('.', ',')


def shortened(text, length=27):
    return text if len(text) <= length else text[:length-1]+'…'


def app_css():
    return '''<style>
    :root {
        --tempo-teal: #00645f;
        --tempo-deep: #084d49;
        --tempo-coral: #c64128;
        --tempo-yellow: #f3c851;
        --tempo-ink: #203330;
    }
    .stApp {
        background: #f4f7f4;
        color: #203330;
        color-scheme: light;
    }
    h1, h2, h3, p, label, button {
        letter-spacing: 0 !important;
    }
    h2 {
        font-size: 1.45rem !important;
        line-height: 1.35 !important;
    }
    h3 {
        font-size: 1.12rem !important;
        line-height: 1.4 !important;
    }
    [data-testid="stMainBlockContainer"] {
        max-width: 1530px;
        padding-top: 3.8rem;
        padding-bottom: 3rem;
    }
    [data-testid="stSidebar"] {
        background: #e5eee8;
        border-right: 1px solid #cad9cf;
    }
    [data-testid="stSidebar"] h2,
    [data-testid="stSidebar"] h3,
    [data-testid="stSidebar"] label,
    [data-testid="stSidebar"] p {
        color: #203330;
    }
    [data-testid="stSidebar"] [data-testid="stBaseButton-tertiary"],
    [data-testid="stSidebar"] [data-testid="stBaseButton-secondary"] {
        background: transparent;
        border: 1px solid transparent;
        color: #25453e;
        justify-content: flex-start;
        min-height: 43px;
    }
    [data-testid="stSidebar"] [data-testid="stBaseButton-tertiary"]:hover,
    [data-testid="stSidebar"] [data-testid="stBaseButton-secondary"]:hover {
        background: #d4e3d9;
        border-color: #b6ccc0;
    }
    [data-testid="stSidebar"] [data-testid="stBaseButton-primary"] {
        background: #00645f;
        border-color: #00645f;
        color: white;
        justify-content: flex-start;
        min-height: 43px;
    }
    [data-testid="stSidebar"] button p {
        color: inherit !important;
    }
    [data-testid="stSidebar"] [data-testid="stExpander"] summary {
        background: #dbe8df;
        color: #203330;
    }
    [data-testid="stCaptionContainer"],
    [data-testid="stCaptionContainer"] p {
        color: #516761 !important;
        opacity: 1 !important;
    }
    [data-testid="stBaseButton-primary"],
    [data-testid="stBaseButton-primaryFormSubmit"] {
        background: #c64128;
        border-color: #c64128;
        color: white;
    }
    [data-testid="stBaseButton-primary"]:hover,
    [data-testid="stBaseButton-primaryFormSubmit"]:hover {
        background: #a93420;
        border-color: #a93420;
        color: white;
    }
    [data-testid="stBaseButton-secondary"],
    [data-testid="stBaseButton-secondaryFormSubmit"] {
        background: white;
        color: #244b42;
        border-color: #b5c7bd;
    }
    button {
        border-radius: 5px !important;
        min-height: 40px;
    }
    button p {
        white-space: normal !important;
        overflow-wrap: anywhere;
    }
    button:focus-visible {
        outline: 3px solid #00645f !important;
        outline-offset: 2px;
    }
    button:disabled {
        opacity: .45;
    }
    [data-testid="stForm"] {
        background: #fff;
        padding: 16px;
        border: 1px solid #c4d3c9;
        border-top: 3px solid #00645f;
        border-radius: 6px;
    }
    [data-testid="stExpander"] [data-testid="stForm"] {
        background: transparent;
        border: 0;
        border-radius: 0;
        padding: 0;
    }
    [data-testid="stMetricValue"] {
        font-size: 1.75rem;
        overflow-wrap: anywhere;
    }
    .tempo-brand {
        display: flex;
        flex-wrap: wrap;
        gap: 20px;
        align-items: center;
        justify-content: space-between;
        padding: 25px 28px;
        color: white;
        background: #00645f;
        border-bottom: 7px solid #f3c851;
        margin-bottom: 18px;
    }
    .tempo-brand-lockup {
        display: flex;
        flex: 1 1 320px;
        gap: 15px;
        align-items: center;
        min-width: 0;
    }
    .tempo-brand-mark {
        display: grid;
        place-items: center;
        width: 50px;
        height: 54px;
        flex: 0 0 50px;
        background: #f3c851;
        color: #203330;
        font-size: 25px;
        font-weight: 850;
        border-radius: 3px;
    }
    .tempo-brand h1 {
        font-size: 32px;
        line-height: 1.2;
        color: white;
        padding: 0;
        margin: 0;
        overflow-wrap: anywhere;
    }
    .tempo-brand p {
        font-size: 13px;
        color: #e0f0e9;
        margin: 6px 0 0;
    }
    .tempo-brand-meta {
        text-align: right;
        max-width: 34%;
        overflow-wrap: anywhere;
    }
    .tempo-brand-meta strong {
        display: block;
        font-size: 14px;
    }
    .tempo-brand-meta span {
        display: block;
        font-size: 12px;
        color: #e0f0e9;
        margin-top: 6px;
    }
    .tempo-sidebrand {
        padding: 10px 2px 22px;
        font-weight: 800;
        font-size: 21px;
        color: #00645f;
    }
    .tempo-sidebrand small {
        display: block;
        color: #516761;
        font-size: 11px;
        font-weight: 500;
        margin-top: 5px;
    }
    .tempo-stats {
        display: grid;
        grid-template-columns: repeat(4, minmax(0, 1fr));
        gap: 12px;
        margin: 10px 0 24px;
    }
    .tempo-stat {
        min-width: 0;
        padding: 16px 18px;
        background: var(--tile-bg);
        color: #203330;
        border-radius: 6px;
        border-bottom: 3px solid var(--tile-edge);
    }
    .tempo-stat span {
        font-size: 12px;
        font-weight: 600;
        display: block;
    }
    .tempo-stat strong {
        display: block;
        font-size: 29px;
        line-height: 1.3;
        padding: 8px 0 3px;
        font-variant-numeric: tabular-nums;
        overflow-wrap: anywhere;
    }
    .tempo-stat small {
        font-size: 12px;
        display: block;
        overflow-wrap: anywhere;
    }
    .tempo-preview {
        display: flex;
        align-items: center;
        gap: 10px;
        background: #fff0c3;
        border-left: 4px solid #a97400;
        padding: 10px 14px;
        margin-bottom: 14px;
        font-size: 13px;
        color: #5b470b;
        overflow-wrap: anywhere;
    }
    .tempo-path {
        padding: 16px 0;
        display: flex;
        gap: 9px;
        flex-wrap: wrap;
        align-items: center;
        border-bottom: 1px solid #c7d5cb;
        margin-bottom: 16px;
    }
    .tempo-path b {
        padding: 6px 10px;
        background: #c64128;
        color: white;
        border-radius: 3px;
        font-size: 13px;
    }
    .tempo-path span {
        color: #7c5044;
        font-size: 17px;
    }
    .tempo-insight {
        padding: 15px 0;
        border-bottom: 1px solid #c7d5cb;
    }
    .tempo-insight strong {
        color: #00645f;
        font-size: 24px;
    }
    .tempo-insight p {
        color: #516761;
        font-size: 13px;
        margin: 5px 0 0;
    }
    .tempo-risk-meter {
        height: 14px;
        background: #dae7df;
        border-radius: 3px;
        overflow: hidden;
        margin: 8px 0 12px;
    }
    .tempo-risk-meter div {
        height: 100%;
        background: #c64128;
    }
    .tempo-phase-legend {
        display: flex;
        gap: 16px;
        flex-wrap: wrap;
        margin: 8px 0 16px;
        font-size: 12px;
    }
    .tempo-phase-legend i {
        display: inline-block;
        width: 10px;
        height: 10px;
        border-radius: 2px;
        margin-right: 5px;
    }
    .tempo-task-status {
        display: flex;
        gap: 12px;
        flex-wrap: wrap;
        margin-bottom: 12px;
        color: #516761;
        font-size: 12px;
    }
    .tempo-footer {
        margin-top: 26px;
        padding: 15px 0;
        border-top: 1px solid #c7d5cb;
        color: #516761;
        font-size: 12px;
    }
    @media (max-width: 1100px) {
        .tempo-stats {
            grid-template-columns: repeat(2, minmax(0, 1fr));
        }
        .tempo-brand {
            padding: 20px;
        }
    }
    @media (max-width: 640px) {
        [data-testid="stMainBlockContainer"] {
            padding: 3.7rem .9rem 2rem;
        }
        .tempo-brand {
            flex-direction: column;
            align-items: flex-start;
            padding: 17px 14px;
            gap: 14px;
        }
        .tempo-brand h1 {
            font-size: 25px;
        }
        .tempo-brand-lockup {
            gap: 10px;
            width: 100%;
            flex: 0 1 auto;
        }
        .tempo-brand-mark {
            width: 36px;
            height: 43px;
            flex-basis: 36px;
            font-size: 22px;
        }
        .tempo-brand-meta {
            max-width: 100%;
            text-align: left;
        }
        .tempo-stats {
            gap: 8px;
            margin-bottom: 18px;
        }
        .tempo-stat {
            padding: 12px 10px;
        }
        .tempo-stat strong {
            font-size: 23px;
        }
        [data-testid="stColumn"] {
            min-width: 0;
            flex: 1 1 100%;
        }
        input {
            font-size: 16px !important;
        }
    }
    @media (max-width: 350px) {
        .tempo-brand h1 {
            font-size: 23px;
        }
    }
    @media (prefers-reduced-motion: reduce) {
        * {
            transition: none !important;
            scroll-behavior: auto !important;
        }
    }
    </style>'''


DIAGRAM = st.components.v2.component(
    'tempo_planning_diagram',
    html='<div id="tempo-drawing"></div>',
    css='''
    :host { display:block; max-width:100%; }
    #tempo-drawing { max-width:100%; }
    .tempo-scroll {
        overflow:auto;
        max-width:100%;
        max-height:730px;
        background:white;
        border:1px solid #c6d7cd;
        border-radius:6px;
    }
    svg { display:block; width:100%; height:auto; }
    svg text { font-family:Arial,sans-serif; letter-spacing:0; }
    .tempo-chart-key { fill:#516761; font-size:12px; }
    .tempo-bar:hover { opacity:.8; }
    @media (prefers-reduced-motion:no-preference) {
        .tempo-bar { transition:opacity .15s ease; }
    }
    ''',
    js='''
    export default function({data,parentElement}) {
        const root=parentElement.querySelector('#tempo-drawing');
        root.innerHTML=data.html;
    }
    ''',
)


def draw_diagram(html, key):
    return DIAGRAM(
        data={'html': html},
        key=key,
        height='content',
        width='stretch',
    )


def gantt_svg(schedule, color_mode='critique', phase=None, critical_only=False):
    frame = schedule['table'].copy()
    if phase and phase != 'Toutes':
        frame = frame.loc[frame['phase'].eq(phase)]
    if critical_only:
        frame = frame.loc[frame['critique']]
    frame = frame.sort_values(['debut_j','id'])
    if frame.empty:
        return '<p style="color:#516761;font:14px Arial">Aucune tâche dans ce filtre.</p>'
    duration = schedule['duree_j']
    xmax = max(duration, 1.)
    width, left, right, row_height = 1020, 252, 58, 49
    plot_width = width-left-right
    height = 74+len(frame)*row_height+35
    scale = plot_width/xmax
    ticks = np.linspace(0, xmax, min(11, max(3, int(math.ceil(xmax))+1)))
    pieces = [
        '<div class="tempo-scroll" tabindex="0" role="region" aria-label="Planning de Gantt, défilement horizontal disponible">',
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" style="min-width:760px">',
        '<title>Planning des tâches, début et fin au plus tôt en jours de travail abstraits</title>',
        f'<rect width="{width}" height="{height}" fill="#ffffff"/>',
        '<text x="19" y="30" fill="#00645f" font-size="13" font-weight="700">TÂCHE / LOT</text>',
        f'<text x="{left}" y="30" fill="#00645f" font-size="13" font-weight="700">JOURS DE TRAVAIL</text>',
        f'<line x1="0" y1="57" x2="{width}" y2="57" stroke="#c6d7cd"/>',
    ]
    for value in ticks:
        x = left+value*scale
        pieces.extend([
            f'<line x1="{x:.3f}" y1="58" x2="{x:.3f}" y2="{height-24}" stroke="#e0e9e3"/>',
            f'<text x="{x:.3f}" y="49" text-anchor="middle" class="tempo-chart-key">J{fmt(value,1)}</text>',
        ])
    for index, row in enumerate(frame.itertuples()):
        top = 69+index*row_height
        if index % 2 == 1:
            pieces.append(f'<rect x="0" y="{top-8}" width="{width}" height="{row_height}" fill="#f5f8f5"/>')
        color = CORAL if row.critique and color_mode == 'critique' else PHASES[row.phase]
        if color_mode == 'critique' and not row.critique:
            color = '#58978d'
        badge_color = '#b83b24' if row.critique and color_mode == 'critique' else DEEP
        label = escape(shortened(row.nom, 26))
        pieces.extend([
            f'<rect x="17" y="{top+2}" width="31" height="27" rx="3" fill="{badge_color}"/>',
            f'<text x="32.5" y="{top+20}" text-anchor="middle" fill="white" font-size="11" font-weight="700">{escape(shortened(row.id,4))}</text>',
            f'<text x="58" y="{top+13}" fill="#203330" font-size="12.5" font-weight="600">{label}</text>',
            f'<text x="58" y="{top+29}" fill="#516761" font-size="10.5">{escape(row.phase)}</text>',
        ])
        x = left+row.debut_j*scale
        bar_width = row.duree_j*scale
        tooltip = escape(
            f'{row.id} — {row.nom} | J{fmt(row.debut_j)} à J{fmt(row.fin_j)} | '
            f'durée {fmt(row.duree_j)} j | marge {fmt(row.marge_j)} j'
        )
        if row.duree_j == 0:
            points = f'{x},{top+6} {x+8},{top+16} {x},{top+26} {x-8},{top+16}'
            pieces.append(f'<polygon class="tempo-bar" points="{points}" fill="{color}"><title>{tooltip}</title></polygon>')
        else:
            pieces.append(
                f'<rect class="tempo-bar" x="{x:.3f}" y="{top+4}" width="{max(.7,bar_width):.3f}" '
                f'height="25" rx="4" fill="{color}"><title>{tooltip}</title></rect>'
            )
        if row.marge_j > 0:
            slack_width = row.marge_j*scale
            pieces.append(
                f'<rect x="{left+row.fin_j*scale:.3f}" y="{top+12}" width="{slack_width:.3f}" '
                f'height="9" fill="#e1ebe4" stroke="#9fb8a9" stroke-dasharray="4,3"><title>Marge totale {fmt(row.marge_j)} j</title></rect>'
            )
        tx = min(width-10, x+bar_width+7)
        pieces.append(f'<text x="{tx:.3f}" y="{top+22}" fill="#203330" font-size="11.5">{fmt(row.duree_j,1)} j</text>')
    pieces.extend([
        f'<text x="18" y="{height-8}" fill="#516761" font-size="11">Barre = durée ; pointillés = marge totale ; losange = tâche de durée nulle.</text>',
        '</svg></div>',
    ])
    return ''.join(pieces)


def network_svg(schedule, focus=None):
    graph = schedule['graph']
    by_id = schedule['table'].set_index('id').to_dict('index')
    all_nodes = list(schedule['order'])
    visible = set(all_nodes)
    if len(all_nodes) > 20:
        focus = focus if focus in graph else all_nodes[0]
        visible = {focus, *sorted(graph.predecessors(focus))[:10], *sorted(graph.successors(focus))[:10]}
    subgraph = graph.subgraph(visible).copy()
    generations = list(nx.topological_generations(subgraph))
    max_rows = max(len(level) for level in generations)
    width = max(700, len(generations)*205+40)
    height = max(270, max_rows*105+95)
    positions = {}
    for column, nodes in enumerate(generations):
        for index, node in enumerate(sorted(nodes)):
            y = 65+(height-110)*(index+.5)/len(nodes)
            positions[node] = (28+205*column, y)
    parts = [
        '<div class="tempo-scroll" tabindex="0" role="region" aria-label="Réseau des dépendances, défilement horizontal disponible">',
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" style="min-width:{min(width,1000)}px">',
        '<title>Réseau fin-début : une tâche attend la fin de tous ses prédécesseurs</title>',
        '<defs><marker id="tempo-arrow" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="6" markerHeight="6" orient="auto-start-reverse"><path d="M0 0L10 5L0 10z" fill="#7e9a8c"/></marker></defs>',
        f'<rect width="{width}" height="{height}" fill="#f8faf7"/>',
        '<text x="22" y="29" fill="#00645f" font-size="14" font-weight="700">LIENS FIN-DÉBUT</text>',
        f'<text x="22" y="49" fill="#516761" font-size="11">{len(visible)} tâches affichées sur {len(all_nodes)}</text>',
    ]
    for start, end in subgraph.edges:
        x1, y1 = positions[start]
        x2, y2 = positions[end]
        parts.append(
            f'<path d="M{x1+160} {y1} C{x1+183} {y1},{x2-23} {y2},{x2-6} {y2}" '
            'fill="none" stroke="#7e9a8c" stroke-width="1.8" marker-end="url(#tempo-arrow)"/>'
        )
    for ident in positions:
        x, y = positions[ident]
        row = by_id[ident]
        color = '#b83b24' if row['critique'] else TEAL
        parts.extend([
            f'<g><title>{escape(ident+" — "+row["nom"])}</title>',
            f'<rect x="{x}" y="{y-33}" width="160" height="68" rx="5" fill="white" stroke="{color}" stroke-width="1.4"/>',
            f'<rect x="{x}" y="{y-33}" width="35" height="68" rx="4" fill="{color}"/>',
            f'<text x="{x+17.5}" y="{y+5}" text-anchor="middle" fill="white" font-size="12" font-weight="700">{escape(shortened(ident,4))}</text>',
            f'<text x="{x+44}" y="{y-10}" fill="#203330" font-size="11" font-weight="700">{escape(shortened(row["nom"],17))}</text>',
            f'<text x="{x+44}" y="{y+8}" fill="#516761" font-size="10.5">J{fmt(row["debut_j"],1)} → J{fmt(row["fin_j"],1)}</text>',
            f'<text x="{x+44}" y="{y+25}" fill="{color}" font-size="10.5">Marge {fmt(row["marge_j"],1)} j</text></g>',
        ])
    parts.append('</svg></div>')
    return ''.join(parts)


def chart(figure, key, height=340):
    figure.update_layout(
        template='plotly_white',
        height=height,
        paper_bgcolor='rgba(0,0,0,0)',
        plot_bgcolor='#ffffff',
        colorway=[TEAL, CORAL, '#8A6BAD', '#B18613'],
        font=dict(family='Arial', size=12, color=INK),
        margin=dict(l=20, r=18, t=28, b=65),
        legend=dict(orientation='h', y=-.23, x=0),
    )
    figure.update_xaxes(automargin=True, gridcolor='#e3ebe5')
    figure.update_yaxes(automargin=True, gridcolor='#e3ebe5')
    st.plotly_chart(
        figure,
        key=key,
        width='stretch',
        config={'displayModeBar':False, 'scrollZoom':False},
    )


def histogram_figure(sim, deadline):
    fig = go.Figure()
    if sim['constant']:
        fig.add_trace(go.Bar(
            x=[sim['moyenne_j']],
            y=[sim['n']],
            marker_color=TEAL,
            width=.2,
            name='Tous les tirages',
        ))
    else:
        fig.add_trace(go.Histogram(
            x=sim['durees'],
            nbinsx=36,
            marker_color=TEAL,
            name='Durées simulées',
            hovertemplate='%{x:.2f} j<br>%{y} tirages<extra></extra>',
        ))
    fig.add_vline(x=deadline, line_dash='dash', line_color=CORAL)
    fig.update_layout(
        xaxis_title='Durée simulée du projet (j)',
        yaxis_title='Nombre de tirages',
        showlegend=False,
    )
    return fig


def cdf_figure(sim, deadline):
    ordered = np.sort(sim['durees'])
    indices = np.unique(np.linspace(0, len(ordered)-1, min(350,len(ordered)), dtype=int))
    fig = go.Figure(go.Scatter(
        x=ordered[indices],
        y=100*(indices+1)/len(ordered),
        mode='lines',
        line=dict(color=CORAL, width=3, shape='hv'),
        name='Cumul des tirages',
        hovertemplate='Durée %{x:.3f} j<br>Cumul %{y:.1f} %<extra></extra>',
    ))
    fig.add_vline(x=deadline, line_dash='dash', line_color=TEAL)
    fig.update_layout(
        xaxis_title='Échéance (j)',
        yaxis_title='Part terminée avant ou à cette durée (%)',
        yaxis_range=[0,100],
        showlegend=False,
    )
    return fig


def criticality_figure(sim):
    frame = sim['influence'].sort_values('criticite_pct', ascending=True)
    labels = frame.apply(lambda row:row['id']+' · '+shortened(row['nom'],22), axis=1)
    fig = go.Figure(go.Bar(
        x=frame['criticite_pct'],
        y=labels,
        orientation='h',
        marker_color=[CORAL if value >= 50 else TEAL for value in frame['criticite_pct']],
        hovertemplate='%{y}<br>Criticité %{x:.2f} %<extra></extra>',
    ))
    fig.update_layout(
        xaxis_title='Part des tirages où la tâche est critique (%)',
        xaxis_range=[0,100],
        showlegend=False,
    )
    return fig


def correlation_figure(sim):
    frame = sim['influence'].dropna(subset=['correlation_rang']).sort_values('correlation_rang')
    fig = go.Figure()
    if not frame.empty:
        labels = frame.apply(lambda row:row['id']+' · '+shortened(row['nom'],22), axis=1)
        fig.add_trace(go.Bar(
            x=frame['correlation_rang'],
            y=labels,
            orientation='h',
            marker_color='#8A6BAD',
            hovertemplate='%{y}<br>Spearman %{x:.3f}<extra></extra>',
        ))
    fig.update_layout(
        xaxis_title='Corrélation de rang avec la durée du projet',
        xaxis_range=[-1,1],
        showlegend=False,
    )
    return fig


def comparison_figure(comp):
    fig = go.Figure()
    for label, sim, color in [
        ('Actuel', comp['actuel'], TEAL),
        ('Variante', comp['variante'], CORAL),
    ]:
        fig.add_trace(go.Histogram(
            x=sim['durees'],
            nbinsx=35,
            histnorm='probability',
            opacity=.65,
            name=label,
            marker_color=color,
        ))
    fig.update_layout(
        barmode='overlay',
        xaxis_title='Durée du projet (j)',
        yaxis_title='Fréquence',
        yaxis_tickformat='.0%',
    )
    return fig


def pdf_styles():
    font_dir = Path(matplotlib.get_data_path())/'fonts'/'ttf'
    for name, filename in [
        ('Tempo', 'DejaVuSans.ttf'),
        ('TempoBold', 'DejaVuSans-Bold.ttf'),
    ]:
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_dir/filename)))
    return {
        'title': ParagraphStyle(
            'TempoTitle', fontName='TempoBold', fontSize=25,
            leading=31, textColor=colors.HexColor(DEEP), spaceAfter=12,
        ),
        'h2': ParagraphStyle(
            'TempoHeading', fontName='TempoBold', fontSize=13,
            leading=18, textColor=colors.HexColor(TEAL), spaceBefore=14,
            spaceAfter=8, keepWithNext=True,
        ),
        'body': ParagraphStyle(
            'TempoBody', fontName='Tempo', fontSize=9,
            leading=14, textColor=colors.HexColor(INK), spaceAfter=8,
            splitLongWords=True,
        ),
        'cell': ParagraphStyle(
            'TempoCell', fontName='Tempo', fontSize=7.5,
            leading=10.5, textColor=colors.HexColor(INK), splitLongWords=True,
        ),
        'head': ParagraphStyle(
            'TempoHeadCell', fontName='TempoBold', fontSize=7.5,
            leading=10.5, textColor=colors.white, splitLongWords=True,
        ),
    }


def pdf_table(headers, rows, widths, styles):
    data = [[Paragraph(escape(str(value)), styles['head']) for value in headers]]
    for row in rows:
        data.append([
            Paragraph(escape(str(value)).replace('\n', '<br/>'), styles['cell'])
            for value in row
        ])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor(TEAL)),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor(PALE)]),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
        ('TOPPADDING', (0,0), (-1,-1), 7),
        ('BOTTOMPADDING', (0,0), (-1,-1), 7),
        ('LINEBELOW', (0,0), (-1,0), .6, colors.HexColor(DEEP)),
    ]))
    return table


def figure_image(figure, width=499):
    stream = io.BytesIO()
    figure.savefig(stream, format='png', dpi=180, bbox_inches='tight', facecolor='white')
    plt.close(figure)
    stream.seek(0)
    asset = Image(stream)
    asset.drawHeight = width*asset.imageHeight/asset.imageWidth
    asset.drawWidth = width
    return asset


def report_gantt(frame, project_finish):
    figure, axis = plt.subplots(figsize=(9.4, 1.4+.36*len(frame)))
    for position, (_, row) in enumerate(frame.iterrows()):
        color = CORAL if row['critique'] else TEAL
        if row['duree_j'] > 0:
            axis.barh(position, row['duree_j'], left=row['debut_j'], color=color, height=.62)
        else:
            axis.scatter([row['debut_j']], [position], color=color, marker='D', s=26)
        if row['marge_j'] > 0:
            axis.plot(
                [row['fin_j'], row['fin_j']+row['marge_j']],
                [position, position], color='#758B82', linestyle='--', linewidth=1.4,
            )
    axis.set_yticks(range(len(frame)))
    axis.set_yticklabels([f'{row.id} | {shortened(row.nom, 30)}' for row in frame.itertuples()], fontsize=8)
    axis.invert_yaxis()
    axis.set_xlim(0, max(1., project_finish)*1.04)
    axis.set_xlabel('Jours de travail abstraits', fontsize=9)
    axis.grid(axis='x', alpha=.2)
    axis.set_axisbelow(True)
    axis.spines[['top','right']].set_visible(False)
    figure.tight_layout()
    return figure_image(figure)


def report_histogram(simulation, deadline):
    figure, axis = plt.subplots(figsize=(9.4, 3.3))
    if simulation['constant']:
        axis.bar([simulation['moyenne_j']], [simulation['n']], width=.3, color=TEAL)
    else:
        axis.hist(simulation['durees'], bins=35, color=TEAL, edgecolor='white')
    axis.axvline(deadline, color=CORAL, linewidth=2, label=f'Échéance : {fmt(deadline)} j')
    axis.set_xlabel('Durée simulée du projet (j)')
    axis.set_ylabel('Nombre de tirages')
    axis.legend(frameon=False, fontsize=9)
    axis.spines[['top','right']].set_visible(False)
    figure.tight_layout()
    return figure_image(figure)


@st.cache_data(max_entries=3, ttl=1800, show_spinner=False)
def pdf_bytes(configuration, created_at):
    project = validate_project(json.loads(configuration))
    styles = pdf_styles()
    stream = io.BytesIO()
    document = SimpleDocTemplate(
        stream, pagesize=A4, rightMargin=48, leftMargin=48,
        topMargin=56, bottomMargin=49, title=f'{APP_NAME} - {project["titre"]}',
        author=APP_NAME,
    )
    story = []

    def paragraph(text, kind='body'):
        return Paragraph(escape(str(text)), styles[kind])

    story.extend([
        paragraph(APP_NAME, 'title'),
        paragraph('Planning, chemin critique et risque de retard', 'h2'),
        paragraph(project['titre']),
        paragraph(f'Origine déclarée : {project["origine"]}. Rapport généré le {created_at}.'),
        paragraph(LIMITS),
        paragraph(
            'Ce rapport décrit le projet enregistré, pas un aperçu. Les résultats '
            'dépendent des estimations saisies et ne constituent pas un engagement de délai.'
        ),
    ])
    if not project['taches']:
        story.append(paragraph('Aucune tâche enregistrée. Aucun calcul de planning effectué.'))
    else:
        schedule = cached_schedule(configuration)
        basis = 'durée la plus probable' if project['parametres']['base'] == 'mode' else 'moyenne triangulaire'
        story.extend([
            paragraph('1. Synthèse du planning', 'h2'),
            pdf_table(
                ['Durée nominale', 'Échéance', 'Tâches', 'Tâches critiques'],
                [[f'{fmt(schedule["duree_j"])} j', f'{fmt(project["parametres"]["echeance_j"])} j',
                  len(project['taches']), schedule['n_critiques']]],
                [125,125,124,125], styles,
            ),
            Spacer(1,10),
            paragraph(f'Base nominale : {basis}. Un chemin critique : {" → ".join(schedule["chemin"])}.'),
            paragraph('Plusieurs chemins critiques sont possibles. La somme des durées ne représente pas la durée totale lorsque des tâches se déroulent en parallèle.'),
            paragraph('2. Planning de Gantt', 'h2'),
            paragraph('Corail : tâche critique. Vert : autre tâche. Tirets : marge totale. Origine du planning : jour 0.'),
        ])
        frame = schedule['table']
        for start in range(0, len(frame), 10):
            story.append(report_gantt(frame.iloc[start:start+10], schedule['duree_j']))
            story.append(Spacer(1,10))
        story.extend([
            paragraph('3. Saisies et calculs vérifiables', 'h2'),
            paragraph('a = durée minimale ; m = durée la plus probable ; b = durée maximale. Toutes les durées et les marges ci-dessous sont en jours.'),
            pdf_table(
                ['ID / tâche', 'Avant', 'a / m / b', 'Début', 'Fin', 'Marge'],
                [[f'{r.id}\n{r.nom}', ', '.join(r.predecesseurs) or 'Aucun',
                  f'{fmt(r.min_j)} / {fmt(r.mode_j)} / {fmt(r.max_j)}',
                  fmt(r.debut_j), fmt(r.fin_j), fmt(r.marge_j)] for r in frame.itertuples()],
                [161,62,105,57,57,57], styles,
            ),
            Spacer(1,10),
            paragraph('Début au plus tôt = plus grande fin des prédécesseurs. Fin = début + durée. Marge totale = début au plus tard - début au plus tôt.'),
        ])
    if project['simulation_active']:
        simulation = cached_simulation(configuration)
        story.extend([
            paragraph('4. Risque de retard simulé', 'h2'),
            paragraph(
                f'{simulation["n"]} tirages, graine {project["parametres"]["graine"]}. '
                'Durées indépendantes suivant des lois triangulaires ; une durée fixe reste constante. '
                'Retard = durée strictement supérieure à l’échéance.'
            ),
            pdf_table(
                ['Moyenne', 'P50', 'P80', 'P95', 'P(retard)'],
                [[f'{fmt(simulation["moyenne_j"])} j',
                  *[f'{fmt(simulation["quantiles"][key])} j' for key in ('P50','P80','P95')],
                  f'{fmt(100*simulation["probabilite"],1)} %']],
                [100,100,100,100,99], styles,
            ),
            Spacer(1,9),
            report_histogram(simulation, project['parametres']['echeance_j']),
        ])
        if not simulation['constant']:
            story.append(paragraph(
                f'Intervalle binomial exact à 95 % : [{fmt(100*simulation["ic_bas"],2)} % ; '
                f'{fmt(100*simulation["ic_haut"],2)} %]. Il décrit l’incertitude numérique '
                'sur la proportion simulée, pas l’incertitude sur la qualité du modèle.'
            ))
        else:
            story.append(paragraph('La durée du projet est constante dans les bornes saisies. Aucun intervalle Monte-Carlo n’est nécessaire.'))
        story.extend([
            paragraph('Tâches à surveiller', 'h2'),
            pdf_table(
                ['ID / tâche', 'Criticité simulée', 'Corrélation de rang'],
                [[f'{r.id} / {r.nom}', f'{fmt(r.criticite_pct,1)} %',
                  fmt(r.correlation_rang,3)] for r in simulation['influence'].itertuples()],
                [285,107,107], styles,
            ),
            Spacer(1,8),
            paragraph('La criticité est la part des tirages où la marge est nulle. La corrélation de rang mesure une association, pas une causalité ; elle est non définie pour une variable constante.'),
        ])
    if project['comparaison_active']:
        comparison = cached_comparison(configuration)
        variant = project['variante']
        story.extend([
            paragraph('5. Comparaison d’une variante', 'h2'),
            paragraph(
                f'Tâche modifiée : {variant["id"]}. Nouvelle estimation a / m / b : '
                f'{fmt(variant["min_j"])} / {fmt(variant["mode_j"])} / {fmt(variant["max_j"])} j. '
                'Les autres tâches et les dépendances restent inchangées.'
            ),
            pdf_table(
                list(comparison['table'].columns),
                [[row[0],*[fmt(value,3) for value in row[1:]]]
                 for row in comparison['table'].itertuples(index=False,name=None)],
                [175,108,108,108], styles,
            ),
            Spacer(1,8),
            paragraph('Les deux simulations utilisent les mêmes nombres uniformes par tâche pour faciliter la comparaison. Ni le coût ni la faisabilité de la réduction ne sont évalués.'),
        ])
    story.extend([
        paragraph('Hypothèses et limites', 'h2'),
        paragraph('Jours abstraits, sans calendrier ni contraintes de ressources. Les tirages ne sont pas des observations. P80 est un quantile conditionnel aux estimations, pas une garantie. Sauvegarder le JSON pour reproduire les paramètres.'),
        paragraph('Références des méthodes', 'h2'),
        paragraph('NetworkX (chemins critiques) : networkx.org/documentation/stable/reference/algorithms/dag.html. SciPy (loi triangulaire, Spearman) : docs.scipy.org/doc/scipy/reference/stats.html.'),
    ])

    def page_footer(canvas, doc):
        canvas.saveState()
        width, height = A4
        canvas.setFillColor(colors.HexColor(TEAL))
        canvas.rect(0, height-12, width, 12, fill=1, stroke=0)
        canvas.setFont('Tempo',7)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(48,28,'TempoChantier | Rapport pédagogique | Jours abstraits')
        canvas.drawRightString(width-48,28,f'Page {doc.page}')
        canvas.restoreState()

    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return stream.getvalue()


def edited_origin(project):
    if project['origine'] in ('Exemple fictif','Exemple fictif modifié'):
        return 'Exemple fictif modifié'
    return project['origine']


def commit_project(project, message):
    clean = validate_project(project)
    st.session_state['tempo_project'] = clean
    st.session_state['tempo_revision'] = st.session_state.get('tempo_revision',0)+1
    st.session_state['tempo_flash'] = message
    st.session_state.pop('tempo_report',None)
    st.rerun()


def navigate(section):
    st.session_state['tempo_section'] = section
    st.rerun()


def show_error(error):
    st.error(str(error), icon=':material/error:')


def table_view(frame, key=None, height='auto', column_config=None):
    st.dataframe(
        frame, hide_index=True, width='stretch', height=height,
        column_config=column_config, key=key,
    )


def stat_tiles(items):
    tiles = []
    palettes = [('#e1f0e9',TEAL),('#fbe8e0',CORAL),('#fcf0c9','#85620b'),('#ece7f4','#715297')]
    for index, (label,value,note) in enumerate(items):
        background, edge = palettes[index%len(palettes)]
        tiles.append(
            f'<div class="tempo-stat" style="--tile-bg:{background};--tile-edge:{edge}">'
            f'<span>{escape(label)}</span><strong>{escape(str(value))}</strong>'
            f'<small>{escape(note)}</small></div>'
        )
    st.html('<div class="tempo-stats">'+''.join(tiles)+'</div>')


def path_badges(path):
    st.html(
        '<div class="tempo-path">'+
        '<span aria-hidden="true">→</span>'.join(f'<b>{escape(ident)}</b>' for ident in path)+
        '</div>'
    )


def require_tasks(project):
    if project['taches']:
        return True
    st.info('Aucune tâche enregistrée dans ce projet.', icon=':material/edit_note:')
    if st.button('Saisir mes tâches', icon=':material/add_task:', key='empty_edit', type='primary'):
        navigate('Tâches et liens')
    return False


def example_picker(project, key):
    choices = {
        'Petit chantier : huit tâches':'chantier',
        'Contrôle simple : sept jours':'sept_jours',
        'Deux chemins concurrents':'deux_chemins',
    }
    with st.form(f'example_form_{key}', border=False):
        choice = st.selectbox('Jeu fictif', list(choices), key=f'example_choice_{key}')
        replace = True
        if project['taches']:
            replace = st.checkbox(
                'Remplacer mes tâches et mes paramètres par cet exemple',
                key=f'example_replace_{key}',
            )
        submitted = st.form_submit_button('Charger cet exemple', icon=':material/science:', width='stretch')
    if submitted:
        if not replace:
            st.warning('Confirmer le remplacement avant de charger un exemple.')
        else:
            commit_project(example_project(choices[choice]),'Exemple fictif chargé. Les anciennes données ont été remplacées.')


def sidebar(project):
    with st.sidebar:
        st.html('<div class="tempo-sidebrand"><strong>TempoChantier</strong><small>Planning & incertitudes</small></div>')
        section = st.session_state.get('tempo_section','Planning')
        for label, icon in SECTIONS:
            if st.button(
                label, icon=f':material/{icon}:', width='stretch',
                type='primary' if section == label else 'tertiary', key=f'nav_{icon}',
            ):
                navigate(label)
        st.divider()
        st.caption('PROJET EN SESSION')
        st.write(project['titre'])
        st.caption(f'{len(project["taches"])} tâche(s) · {project["origine"]}')
        st.caption('Les données de cette session ne sont pas enregistrées sur disque.')
        with st.expander('Exemples fictifs', icon=':material/science:'):
            example_picker(project,'sidebar')
        with st.expander('Nouveau projet', icon=':material/create_new_folder:'):
            confirmed = st.checkbox('Effacer le projet en session', key='confirm_empty')
            if st.button('Créer un projet vide', icon=':material/add:', width='stretch'):
                if confirmed:
                    commit_project(new_project(),'Nouveau projet vide créé.')
                else:
                    st.warning('Confirmer l’effacement du projet en session.')
        st.divider()
        st.caption('Atelier pédagogique · Génie civil')
        st.caption('Unité : jour de travail abstrait')


def brand(project):
    st.html(
        '<header class="tempo-brand"><div class="tempo-brand-lockup">'
        '<div class="tempo-brand-mark" aria-hidden="true">T</div><div>'
        '<h1>TempoChantier</h1><p>Planning · Chemin critique · Risque de retard</p>'
        '</div></div><div class="tempo-brand-meta">'
        f'<strong>{escape(project["titre"])}</strong><span>{escape(project["origine"])}</span>'
        '</div></header>'
    )


def planning_view(project):
    preview = not project['taches']
    displayed = example_project() if preview else project
    result = cached_schedule(canonical(displayed))
    if preview:
        st.html(
            '<div class="tempo-preview"><strong>Aperçu fictif</strong>'
            '<span>Aucune tâche utilisateur enregistrée. Cet aperçu ne fait pas partie de votre projet.</span></div>'
        )
        left, right = st.columns(2)
        if left.button('Charger le chantier fictif', icon=':material/science:', type='primary', width='stretch', wrap=True):
            commit_project(example_project(),'Chantier fictif chargé : huit tâches, aucune observation réelle.')
        if right.button('Créer mon planning', icon=':material/add_task:', width='stretch', wrap=True):
            navigate('Tâches et liens')
    stat_tiles([
        ('Durée nominale',f'{fmt(result["duree_j"])} j','Sans contrainte de ressources'),
        ('Échéance',f'{fmt(displayed["parametres"]["echeance_j"])} j','Seuil de comparaison'),
        ('Tâches critiques',f'{result["n_critiques"]} / {len(displayed["taches"])}','Marge totale nulle'),
        ('Somme des durées',f'{fmt(result["somme_durees_j"])} j','Différente de la durée du projet'),
    ])
    st.subheader('Planning de Gantt')
    a,b,c = st.columns([1.1,1.1,1])
    mode = a.selectbox('Couleurs', ['Chemin critique','Phases'], key='gantt_color')
    phase = b.selectbox('Phase affichée', ['Toutes',*PHASES], key='gantt_phase')
    only = c.toggle('Tâches critiques', key='gantt_only_critical', wrap=True)
    draw_diagram(
        gantt_svg(result, 'phase' if mode == 'Phases' else 'critique',
                  None if phase == 'Toutes' else phase, only),
        'planning_gantt',
    )
    st.caption('Jour 0 = début du planning. Corail : tâche critique. Tirets : marge totale. Les filtres ne changent pas les calculs.')
    if mode == 'Phases':
        st.html('<div class="tempo-phase-legend">'+''.join(
            f'<span><i style="background:{color}"></i>{escape(label)}</span>' for label,color in PHASES.items()
        )+'</div>')
    st.subheader('Un chemin qui fixe la durée')
    path_badges(result['chemin'])
    st.caption('Plusieurs chemins critiques peuvent coexister. Réduire une tâche non critique ne réduit pas nécessairement la durée du projet.')
    delta = displayed['parametres']['echeance_j']-result['duree_j']
    if delta >= 0:
        st.success(f'La durée nominale laisse {fmt(delta)} j avant l’échéance. Ce n’est pas une garantie face aux aléas.', icon=':material/event_available:')
    else:
        st.warning(f'La durée nominale dépasse déjà l’échéance de {fmt(-delta)} j.', icon=':material/event_busy:')
    st.subheader('Zoom sur une tâche')
    ids = result['order']
    names = {row['id']:row['nom'] for row in displayed['taches']}
    selected = st.selectbox('Tâche',ids,format_func=lambda value:f'{value} · {names[value]}',key='planning_focus')
    detail = result['table'].set_index('id').loc[selected]
    stat_tiles([
        ('Début au plus tôt',f'{fmt(detail["debut_j"])} j','Après tous les prédécesseurs'),
        ('Fin au plus tôt',f'{fmt(detail["fin_j"])} j','Début + durée retenue'),
        ('Marge totale',f'{fmt(detail["marge_j"])} j','Sans repousser la fin du projet'),
        ('Marge libre',f'{fmt(detail["marge_libre_j"])} j','Sans repousser le début des suivantes'),
    ])
    draw_diagram(network_svg(result,selected),'planning_network')
    st.caption('Les flèches sont des dépendances fin-début, pas un trajet géographique. Pour plus de 20 tâches : voisinage de la tâche sélectionnée.')
    with st.expander('Tableau des dates et marges', icon=':material/table_chart:'):
        frame = result['table'][['id','nom','duree_j','debut_j','fin_j','debut_tard_j','fin_tard_j','marge_j','marge_libre_j','critique']]
        table_view(frame)
    with st.expander('Travaux en parallèle', icon=':material/stacks:'):
        concurrency = concurrency_table(result)
        if concurrency.empty:
            st.info('Toutes les durées sont nulles : aucun intervalle de travail positif.')
        else:
            maximum = int(concurrency['taches_simultanees'].max())
            st.write(f'Au maximum, **{maximum} tâches** sont actives en même temps dans ce planning au plus tôt.')
            table_view(concurrency)
            st.caption('Ce comptage ne calcule ni le nombre d’ouvriers ni les besoins en matériel.')
    if not preview:
        with st.expander('Réglages du planning', icon=':material/tune:'):
            planning_settings(project)


def planning_settings(project):
    revision = st.session_state['tempo_revision']
    with st.form(f'planning_settings_{revision}', border=False):
        title = st.text_input('Nom du projet',project['titre'],max_chars=100)
        basis = st.radio(
            'Durée utilisée pour le planning nominal', ['mode','moyenne'],
            index=0 if project['parametres']['base'] == 'mode' else 1,
            format_func=lambda value:'La plus probable (m)' if value == 'mode' else 'Moyenne triangulaire (a + m + b) / 3',
        )
        deadline = st.number_input('Échéance (j)',min_value=0.,max_value=20000.,value=project['parametres']['echeance_j'],step=.5)
        submitted = st.form_submit_button('Appliquer les réglages', icon=':material/check:', type='primary')
    if submitted:
        candidate = copy.deepcopy(project)
        candidate['titre'] = title
        candidate['parametres'].update(base=basis,echeance_j=deadline)
        candidate.update(simulation_active=False,comparaison_active=False)
        try:
            commit_project(candidate,'Réglages appliqués. Relancer les analyses si nécessaire.')
        except ValueError as error:
            show_error(error)


def task_column_config():
    return {
        'id':st.column_config.TextColumn('ID',required=True,max_chars=16,width='small'),
        'nom':st.column_config.TextColumn('Nom de la tâche',required=True,max_chars=80,width='large'),
        'phase':st.column_config.SelectboxColumn('Phase',options=list(PHASES),required=True),
        'min_j':st.column_config.NumberColumn('Min. a (j)',min_value=0.,max_value=365.,step=.5,required=True),
        'mode_j':st.column_config.NumberColumn('Probable m (j)',min_value=0.,max_value=365.,step=.5,required=True),
        'max_j':st.column_config.NumberColumn('Max. b (j)',min_value=0.,max_value=365.,step=.5,required=True),
        'predecesseurs':st.column_config.TextColumn('Avant cette tâche',help='Identifiants séparés par des virgules, par exemple A, B.'),
    }


def next_identifier(project):
    identifiers = {row['id'] for row in project['taches']}
    for index in range(1,MAX_TASKS+2):
        ident = f'T{index}'
        if ident not in identifiers:
            return ident
    raise ValueError('Aucun identifiant disponible.')


def add_task_form(project):
    revision = st.session_state['tempo_revision']
    with st.form(f'add_task_{revision}'):
        a,b = st.columns([1,3])
        ident = a.text_input('Identifiant',value=next_identifier(project),max_chars=16)
        name = b.text_input('Nom de la tâche',max_chars=80,placeholder='Ex. Préparer le coffrage')
        phase = st.selectbox('Phase',list(PHASES))
        a,b,c = st.columns(3)
        low = a.number_input('Minimum a (j)',min_value=0.,max_value=365.,value=1.,step=.5)
        mode = b.number_input('Plus probable m (j)',min_value=0.,max_value=365.,value=2.,step=.5)
        high = c.number_input('Maximum b (j)',min_value=0.,max_value=365.,value=3.,step=.5)
        labels = {row['id']:row['nom'] for row in project['taches']}
        predecessors = st.multiselect(
            'Tâches à terminer avant celle-ci',list(labels),
            format_func=lambda value:f'{value} · {labels[value]}',max_selections=10,
        )
        submitted = st.form_submit_button('Ajouter la tâche',icon=':material/add_task:',type='primary',width='stretch')
    if submitted:
        try:
            row = task(ident,name,phase,low,mode,high)
            row['predecesseurs'] = predecessors
            candidate = candidate_from_tasks(project,[*project['taches'],row],edited_origin(project))
            commit_project(candidate,f'Tâche {ident} ajoutée. Les analyses précédentes ont été réinitialisées.')
        except ValueError as error:
            show_error(error)


def edit_one_task(project):
    revision = st.session_state['tempo_revision']
    names = {row['id']:row['nom'] for row in project['taches']}
    selected = st.selectbox(
        'Tâche à modifier',list(names),format_func=lambda value:f'{value} · {names[value]}',
        key=f'edit_selected_{revision}',
    )
    row = next(row for row in project['taches'] if row['id'] == selected)
    with st.form(f'edit_task_{revision}_{selected}',border=False):
        st.caption(f'ID conservé : {selected}. Le tableau complet permet de modifier les identifiants et leurs liens ensemble.')
        name = st.text_input('Libellé',row['nom'],max_chars=80)
        phase = st.selectbox('Phase de la tâche',list(PHASES),index=list(PHASES).index(row['phase']))
        a,b,c = st.columns(3)
        low = a.number_input('Min. (j)',0.,365.,row['min_j'],.5)
        mode = b.number_input('Probable (j)',0.,365.,row['mode_j'],.5)
        high = c.number_input('Max. (j)',0.,365.,row['max_j'],.5)
        predecessors = st.multiselect(
            'Prédécesseurs', [ident for ident in names if ident != selected],
            default=row['predecesseurs'],format_func=lambda value:f'{value} · {names[value]}',
            max_selections=10,
        )
        submitted = st.form_submit_button('Enregistrer cette tâche',icon=':material/save:',type='primary')
    if submitted:
        updated = task(selected,name,phase,low,mode,high)
        updated['predecesseurs'] = predecessors
        rows = [updated if item['id'] == selected else item for item in project['taches']]
        try:
            commit_project(
                candidate_from_tasks(project,rows,edited_origin(project)),
                f'Tâche {selected} modifiée. Les analyses précédentes ont été réinitialisées.',
            )
        except ValueError as error:
            show_error(error)


def task_deletion_form(project):
    revision = st.session_state['tempo_revision']
    names = {row['id']:row['nom'] for row in project['taches']}
    chosen = st.selectbox(
        'Tâche à supprimer',list(names),format_func=lambda value:f'{value} · {names[value]}',
        key=f'delete_selected_{revision}',
    )
    successors = [row['id'] for row in project['taches'] if chosen in row['predecesseurs']]
    if successors:
        st.warning(f'Cette tâche est un prédécesseur de : {", ".join(successors)}. Retirer ces liens peut avancer ces tâches.')
    with st.form(f'delete_task_{revision}_{chosen}',border=False):
        detach = st.checkbox('Retirer aussi les liens des tâches suivantes',value=False,disabled=not successors)
        confirm = st.checkbox(f'Confirmer la suppression de {chosen}')
        submitted = st.form_submit_button('Supprimer la tâche',icon=':material/delete:')
    if submitted:
        if not confirm:
            st.warning('La suppression doit être confirmée.')
            return
        if successors and not detach:
            st.warning('Supprimer d’abord les liens ou confirmer leur retrait.')
            return
        rows = copy.deepcopy([row for row in project['taches'] if row['id'] != chosen])
        for row in rows:
            row['predecesseurs'] = [value for value in row['predecesseurs'] if value != chosen]
        try:
            commit_project(candidate_from_tasks(project,rows,edited_origin(project)),f'Tâche {chosen} supprimée.')
        except ValueError as error:
            show_error(error)


def task_editor(project):
    revision = st.session_state['tempo_revision']
    st.caption('Toutes les durées sont en jours abstraits. Respecter 0 ≤ a ≤ m ≤ b. Une tâche fixe utilise a = m = b.')
    with st.form(f'bulk_editor_{revision}',border=False):
        edited = st.data_editor(
            tasks_frame(project['taches']),
            column_config=task_column_config(),
            hide_index=True,width='stretch',num_rows='dynamic',
            key=f'task_grid_{revision}',
        )
        submitted = st.form_submit_button('Appliquer le tableau',icon=':material/check:',type='primary')
    if submitted:
        try:
            rows = edited.to_dict('records')
            candidate = candidate_from_tasks(project,rows,edited_origin(project))
            commit_project(candidate,'Tableau validé et enregistré. Aucun cycle détecté.')
        except ValueError as error:
            show_error(error)
            st.caption('Les anciennes données restent conservées tant que le tableau n’est pas valide.')


def tasks_view(project):
    st.header('Tâches et dépendances')
    st.write('Une tâche peut commencer quand **toutes** ses tâches précédentes sont terminées.')
    if len(project['taches']) < MAX_TASKS:
        st.subheader('Ajouter une tâche')
        add_task_form(project)
    else:
        st.info(f'La limite pédagogique de {MAX_TASKS} tâches est atteinte.')
    if project['taches']:
        st.subheader('Projet enregistré')
        table_view(tasks_frame(project['taches']),column_config=task_column_config())
        with st.expander('Modifier une tâche',icon=':material/edit:'):
            edit_one_task(project)
        with st.expander('Modifier le tableau complet',icon=':material/table_edit:'):
            task_editor(project)
        with st.expander('Supprimer une tâche',icon=':material/delete:'):
            task_deletion_form(project)
        with st.expander('Nom et base du planning',icon=':material/tune:'):
            planning_settings(project)
    else:
        st.info('Projet vide : commencer par une tâche sans prédécesseur, ou charger un exemple fictif depuis le menu latéral.')
    st.subheader('Exemple de lien')
    st.html(
        '<div class="tempo-insight"><strong>A → B et A → C</strong>'
        '<p>Après A, B et C peuvent avancer en parallèle. Une tâche D dépendant de B et C attend la fin des deux.</p></div>'
    )
    st.caption('Les boucles A → B → A, les identifiants inconnus et les durées inversées sont refusés. Un identifiant est sensible aux majuscules : A et a sont différents.')


def simulation_settings(project):
    revision = st.session_state['tempo_revision']
    settings = project['parametres']
    with st.form(f'sim_settings_{revision}'):
        a,b,c = st.columns(3)
        deadline = a.number_input('Échéance à tester (j)',0.,20000.,settings['echeance_j'],.5)
        draws = b.number_input('Nombre de tirages',100,50000,settings['tirages'],100)
        seed = c.number_input('Graine aléatoire',0,2147483647,settings['graine'],1)
        submitted = st.form_submit_button('Simuler le risque de retard',icon=':material/play_arrow:',type='primary',width='stretch')
    if submitted:
        candidate = copy.deepcopy(project)
        candidate['parametres'].update(echeance_j=deadline,tirages=draws,graine=seed)
        candidate.update(simulation_active=True,comparaison_active=False)
        try:
            with st.spinner('Simulation des durées et recalcul des chemins…'):
                cached_simulation(canonical(candidate))
            commit_project(candidate,'Simulation terminée. Résultats reproductibles avec la même graine et les mêmes paramètres.')
        except (ValueError,ArithmeticError) as error:
            show_error(error)


def probability_summary(simulation, deadline):
    risk = 100*simulation['probabilite']
    stat_tiles([
        ('Risque de retard',f'{fmt(risk,1)} %',f'Durée strictement > {fmt(deadline)} j'),
        ('Durée moyenne',f'{fmt(simulation["moyenne_j"])} j','Moyenne des tirages du projet'),
        ('P80',f'{fmt(simulation["quantiles"]["P80"])} j','Quantile de la simulation'),
        ('P95',f'{fmt(simulation["quantiles"]["P95"])} j','Quantile de la simulation'),
    ])
    st.html(
        f'<div class="tempo-risk-meter" role="meter" aria-label="Risque simulé de retard" '
        f'aria-valuemin="0" aria-valuemax="100" aria-valuenow="{risk:.4f}">'
        f'<div style="width:{risk:.4f}%"></div></div>'
    )
    st.write(f'**{simulation["depassements"]:,} tirages sur {simulation["n"]:,}** dépassent l’échéance.'.replace(',',' '))
    if simulation['constant']:
        st.info('Durée totale constante dans les bornes saisies : le résultat est déterministe, même si une tâche non critique varie.')
    else:
        st.caption(
            f'Intervalle binomial exact à 95 % : {fmt(100*simulation["ic_bas"])} % à '
            f'{fmt(100*simulation["ic_haut"])} %. Il mesure l’erreur d’échantillonnage, '
            'pas la fiabilité des hypothèses.'
        )


def triangular_figure(row):
    low,mode,high = [row[key] for key in ('min_j','mode_j','max_j')]
    figure = go.Figure()
    if low == high:
        figure.add_trace(go.Scatter(x=[low,low],y=[0,1],mode='lines',line_color=TEAL,name='Durée fixe'))
        figure.update_layout(yaxis_title='Repère de durée fixe',yaxis_showticklabels=False)
    else:
        x = np.linspace(low,high,301)
        density = stats.triang.pdf(x,(mode-low)/(high-low),loc=low,scale=high-low)
        figure.add_trace(go.Scatter(
            x=x,y=density,fill='tozeroy',fillcolor='rgba(0,100,95,.14)',
            line_color=TEAL,mode='lines',name='Densité triangulaire',
        ))
        figure.add_vline(x=mode,line_color=CORAL,line_dash='dash')
        figure.update_layout(yaxis_title='Densité (1/j)')
    figure.update_layout(xaxis_title='Durée de la tâche (j)',showlegend=False)
    return figure


def deadline_comparison(simulation, project):
    st.subheader('Changer de perspective : plusieurs échéances')
    with st.form(f'deadline_compare_{st.session_state["tempo_revision"]}',border=False):
        a,b,c = st.columns(3)
        first = a.number_input('Échéance courte (j)',0.,20000.,max(0.,project['parametres']['echeance_j']-2),.5)
        middle = b.number_input('Échéance centrale (j)',0.,20000.,project['parametres']['echeance_j'],.5)
        last = c.number_input('Échéance longue (j)',0.,20000.,min(20000.,project['parametres']['echeance_j']+2),.5)
        apply = st.form_submit_button('Comparer ces échéances',icon=':material/compare_arrows:')
    key = f'deadlines_{st.session_state["tempo_revision"]}'
    if apply:
        st.session_state[key] = sorted(set([first,middle,last]))
    deadlines = st.session_state.get(key,sorted(set([
        max(0.,project['parametres']['echeance_j']-2),project['parametres']['echeance_j'],
        min(20000.,project['parametres']['echeance_j']+2),
    ])))
    frame = deadline_table(simulation,deadlines)
    table_view(frame)
    st.caption('Les mêmes durées simulées sont réutilisées. Augmenter l’échéance ne peut pas augmenter le risque de dépassement.')


def risk_view(project):
    st.header('Risque de retard')
    if not require_tasks(project):
        return
    st.write('Chaque tâche reçoit une durée plausible entre **a** et **b**, avec **m** comme valeur la plus probable. Le planning complet est recalculé à chaque tirage.')
    simulation_settings(project)
    if not project['simulation_active']:
        st.info('Les résultats apparaîtront après le lancement de la simulation. Modifier les tâches invalide les anciens résultats.')
    else:
        simulation = cached_simulation(canonical(project))
        deadline = project['parametres']['echeance_j']
        probability_summary(simulation,deadline)
        st.subheader('Durées possibles du projet')
        left,right = st.columns(2)
        with left:
            chart(histogram_figure(simulation,deadline),'risk_hist')
        with right:
            chart(cdf_figure(simulation,deadline),'risk_cdf')
        st.caption('Histogramme : fréquence des durées. Courbe cumulée : part des durées inférieures ou égales à une valeur ; affichage allégé à 350 points, calculs réalisés sur tous les tirages.')
        table_view(pd.DataFrame([
            {'Indicateur':'Borne minimale théorique','Durée (j)':simulation['support_min_j']},
            {'Indicateur':'P5 simulé','Durée (j)':simulation['quantiles']['P5']},
            {'Indicateur':'P50 simulé','Durée (j)':simulation['quantiles']['P50']},
            {'Indicateur':'P80 simulé','Durée (j)':simulation['quantiles']['P80']},
            {'Indicateur':'P95 simulé','Durée (j)':simulation['quantiles']['P95']},
            {'Indicateur':'Borne maximale théorique','Durée (j)':simulation['support_max_j']},
        ]))
        st.subheader('Quelles tâches méritent l’attention ?')
        chart(criticality_figure(simulation),'risk_criticality',height=min(1100,max(320,36*len(project['taches'])+120)))
        st.caption('Une tâche critique dans 80 % des tirages a une marge nulle dans 80 % de ces plannings. Ce n’est pas une probabilité de défaillance.')
        variable = simulation['influence']['correlation_rang'].notna().any()
        with st.expander('Association entre les durées et la fin du projet',icon=':material/analytics:'):
            if variable:
                chart(correlation_figure(simulation),'risk_sensitivity',height=min(1100,max(320,36*len(project['taches'])+120)))
            else:
                st.info('Corrélations non définies : durée du projet ou durées des tâches constantes.')
            table_view(simulation['influence'])
            st.caption('Coefficient de Spearman entre la durée d’une tâche et celle du projet. Association de rang, sans preuve de causalité. Une valeur proche de zéro peut masquer un effet local.')
        deadline_comparison(simulation,project)
    with st.expander('Inspecter une estimation de durée',icon=':material/show_chart:'):
        names = {row['id']:row['nom'] for row in project['taches']}
        selected = st.selectbox('Estimation à visualiser',list(names),format_func=lambda value:f'{value} · {names[value]}')
        row = next(row for row in project['taches'] if row['id'] == selected)
        chart(triangular_figure(row),'risk_triangle',height=290)
        mean = (row['min_j']+row['mode_j']+row['max_j'])/3
        st.write(f'Moyenne théorique de cette tâche : **{fmt(mean)} j**. Les trois estimations sont des hypothèses, pas des données observées.')
    st.caption('Hypothèse forte : les durées sont indépendantes. Un aléa météo partagé ou une équipe unique peuvent rendre cette hypothèse inadéquate.')


def scenario_form(project):
    revision = st.session_state['tempo_revision']
    names = {row['id']:row['nom'] for row in project['taches']}
    default = project['variante']['id'] or next(iter(names))
    selected = st.selectbox(
        'Tâche à modifier dans la variante',list(names),index=list(names).index(default),
        format_func=lambda value:f'{value} · {names[value]}',key=f'variant_task_{revision}',
    )
    original = next(row for row in project['taches'] if row['id'] == selected)
    saved = project['variante'] if selected == project['variante']['id'] else original
    st.write(f'Estimation actuelle : **{fmt(original["min_j"])} / {fmt(original["mode_j"])} / {fmt(original["max_j"])} j**.')
    with st.form(f'variant_form_{revision}_{selected}'):
        a,b,c = st.columns(3)
        low = a.number_input('Nouveau minimum (j)',0.,365.,saved['min_j'],.5)
        mode = b.number_input('Nouveau plus probable (j)',0.,365.,saved['mode_j'],.5)
        high = c.number_input('Nouveau maximum (j)',0.,365.,saved['max_j'],.5)
        st.caption(
            f'Comparaison avec {project["parametres"]["tirages"]:,} tirages, '
            f'graine {project["parametres"]["graine"]}, échéance {fmt(project["parametres"]["echeance_j"])} j.'.replace(',',' ')
        )
        submitted = st.form_submit_button('Comparer la variante',icon=':material/compare_arrows:',type='primary',width='stretch')
    if submitted:
        candidate = copy.deepcopy(project)
        candidate['variante'] = {'id':selected,'min_j':low,'mode_j':mode,'max_j':high}
        candidate['comparaison_active'] = True
        try:
            with st.spinner('Comparaison de deux plannings avec tirages appariés…'):
                cached_comparison(canonical(candidate))
            commit_project(candidate,'Comparaison prête. Les tâches du projet actuel restent inchangées.')
        except (ValueError,ArithmeticError) as error:
            show_error(error)


def scenarios_view(project):
    st.header('Comparer une variante')
    if not require_tasks(project):
        return
    st.write('Modifier l’estimation d’**une seule tâche** et mesurer son effet, sans remplacer le projet actuel.')
    scenario_form(project)
    if not project['comparaison_active']:
        st.info('Choisir une tâche et appliquer une nouvelle estimation pour afficher la comparaison.')
        return
    comparison = cached_comparison(canonical(project))
    current = comparison['actuel']
    alternative = comparison['variante']
    delta = comparison['difference_moyenne_j']
    stat_tiles([
        ('Moyenne actuelle',f'{fmt(current["moyenne_j"])} j','Projet enregistré'),
        ('Moyenne variante',f'{fmt(alternative["moyenne_j"])} j',f'Tâche {project["variante"]["id"]} modifiée'),
        ('Écart moyen',f'{fmt(delta)} j','Variante moins actuel'),
        ('Variante plus courte',f'{fmt(100*comparison["part_plus_courte"],1)} %','Part des paires de tirages'),
    ])
    table_view(comparison['table'])
    chart(comparison_figure(comparison),'scenario_distributions',height=380)
    st.caption('Les histogrammes sont normalisés séparément. Une fréquence dépend du découpage en classes ; utiliser le tableau pour comparer les quantiles.')
    if delta < -1e-9:
        st.success(f'La variante réduit la durée moyenne simulée de {fmt(-delta)} j, sous les hypothèses choisies.')
    elif delta > 1e-9:
        st.warning(f'La variante augmente la durée moyenne simulée de {fmt(delta)} j.')
    else:
        st.info('Pas d’écart de durée moyenne détecté avec ces estimations et cette précision.')
    with st.expander('Comparer les chemins nominaux',icon=':material/account_tree:'):
        st.write('**Projet actuel**')
        path_badges(current['nominal']['chemin'])
        st.write('**Variante**')
        path_badges(alternative['nominal']['chemin'])
        draw_diagram(gantt_svg(alternative['nominal']),'scenario_gantt')
    with st.expander('Détail des écarts appariés',icon=':material/bar_chart:'):
        differences = alternative['durees']-current['durees']
        figure = go.Figure(go.Histogram(x=differences,nbinsx=35,marker_color='#8A6BAD'))
        figure.update_layout(xaxis_title='Durée variante - durée actuelle (j)',yaxis_title='Nombre de paires',showlegend=False)
        figure.add_vline(x=0,line_color=TEAL)
        chart(figure,'scenario_paired')
        st.caption('Même nombre uniforme pour une même tâche dans chaque paire. Les paramètres d’une seule tâche changent. Les tirages restent indépendants d’une tâche à l’autre.')
    st.caption('Aucune évaluation des coûts, de la disponibilité des équipes ou de la faisabilité technique. Une réduction de durée ne constitue pas une recommandation de chantier.')


def import_panel(project):
    st.subheader('Importer un projet ou des tâches')
    uploaded = st.file_uploader('Fichier CSV, Excel ou JSON',type=['csv','xlsx','json'],max_upload_size=2)
    if uploaded is None:
        st.caption('CSV / Excel : remplace les tâches et conserve les paramètres. JSON : restaure un projet complet.')
        return
    content = uploaded.getvalue()
    digest = hashlib.sha256(content).hexdigest()[:16]
    try:
        if uploaded.name.lower().endswith('.json'):
            candidate = import_project(content)
            st.write(f'Projet : **{candidate["titre"]}** · {len(candidate["taches"])} tâches · {candidate["origine"]}')
            st.caption('Le JSON restaure aussi l’échéance, la graine et les indicateurs d’analyses actives. Son origine déclarée n’est pas une preuve de provenance.')
        else:
            rows = import_tasks(content,uploaded.name)
            candidate = candidate_from_tasks(project,rows,'Import : origine à documenter')
            st.caption('Les simulations et la variante précédentes seront réinitialisées.')
        table_view(tasks_frame(candidate['taches']))
        with st.form(f'import_confirm_{digest}_{st.session_state["tempo_revision"]}',border=False):
            confirmed = st.checkbox('Remplacer le projet en session par cet import')
            submitted = st.form_submit_button('Valider l’import',icon=':material/upload:',type='primary')
        if submitted:
            if not confirmed:
                st.warning('Confirmer le remplacement avant de poursuivre.')
            else:
                commit_project(candidate,'Import validé. Le fichier original n’a pas été modifié.')
    except (ValueError,TypeError,UnicodeError,zipfile.BadZipFile,OSError) as error:
        show_error(ValueError(f'Import refusé : {error}'))
    except Exception:
        st.error('Ce fichier ne peut pas être lu. Utiliser le modèle CSV ou un classeur Excel .xlsx non chiffré.')


def downloads_panel(project):
    configuration = canonical(project)
    sheets = result_sheets(project)
    st.subheader('Exporter le projet enregistré')
    a,b,c = st.columns(3)
    a.download_button(
        'Projet JSON',configuration.encode('utf-8'),file_name='tempochantier-projet.json',
        mime='application/json',icon=':material/data_object:',width='stretch',key='download_project',wrap=True,
    )
    b.download_button(
        'Tâches CSV',safe_csv(tasks_frame(project['taches'])),file_name='tempochantier-taches.csv',
        mime='text/csv',icon=':material/table_view:',width='stretch',key='download_tasks',wrap=True,
    )
    c.download_button(
        'Classeur Excel',excel_bytes(sheets),file_name='tempochantier-resultats.xlsx',
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        icon=':material/table_chart:',width='stretch',key='download_excel',wrap=True,
    )
    st.caption('JSON : sauvegarde complète et réimportable. Excel : saisies, résultats et analyses actives. CSV : séparateur virgule ; noms des colonnes conservés pour le réimport.')
    if project['taches']:
        frame = sheets['Calculs_CPM']
        st.download_button(
            'Résultats du planning CSV',safe_csv(frame),file_name='tempochantier-planning.csv',
            mime='text/csv',icon=':material/download:',key='download_cpm',
        )
    if project['simulation_active']:
        simulation = cached_simulation(configuration)
        st.download_button(
            'Tous les tirages CSV',safe_csv(simulation['table']),file_name='tempochantier-simulations.csv',
            mime='text/csv',icon=':material/download:',key='download_simulations',
        )
    st.caption('Protection tableur : les textes commençant par =, +, - ou @ sont neutralisés dans les exports CSV. Le JSON conserve le texte exact.')
    st.subheader('Rapport PDF')
    if st.button('Préparer le rapport PDF',icon=':material/picture_as_pdf:',type='primary'):
        try:
            created_at = datetime.now(ZoneInfo('Europe/Paris')).strftime('%d/%m/%Y à %H:%M %Z')
            with st.spinner('Mise en page du rapport…'):
                result = pdf_bytes(configuration,created_at)
            st.session_state['tempo_report'] = {'configuration':configuration,'bytes':result}
        except Exception:
            st.error('Le PDF n’a pas pu être généré. Le projet reste intact ; exporter le JSON ou le classeur Excel.')
    report = st.session_state.get('tempo_report')
    if report and report['configuration'] == configuration:
        st.download_button(
            'Télécharger le rapport PDF',report['bytes'],file_name='tempochantier-rapport.pdf',
            mime='application/pdf',icon=':material/download:',width='stretch',key='download_pdf',
        )
    st.caption('Le rapport porte uniquement sur le projet enregistré. Il n’inclut pas l’aperçu fictif ni les changements non appliqués.')


def data_view(project):
    st.header('Données et rapports')
    st.info('Les données vivent dans la session du navigateur. Exporter le JSON pour les conserver et les restaurer ensuite.',icon=':material/save:')
    downloads_panel(project)
    st.divider()
    import_panel(project)
    with st.expander('Format des fichiers de tâches',icon=':material/help:'):
        table_view(pd.DataFrame([
            {'Colonne':'id','Contenu':'Identifiant unique, sans espace ni accent ; 16 caractères au plus'},
            {'Colonne':'nom','Contenu':'Nom de la tâche ; 80 caractères au plus'},
            {'Colonne':'phase','Contenu':'Préparation, Structure, Réseaux, Finitions ou Autre'},
            {'Colonne':'min_j','Contenu':'Durée minimale a, entre 0 et 365 jours'},
            {'Colonne':'mode_j','Contenu':'Durée la plus probable m, avec a ≤ m ≤ b'},
            {'Colonne':'max_j','Contenu':'Durée maximale b, entre 0 et 365 jours'},
            {'Colonne':'predecesseurs','Contenu':'Identifiants séparés par des virgules ; vide pour commencer au jour 0'},
        ]))
        st.caption('Les colonnes phase et predecesseurs sont facultatives à l’import. Une phase manquante devient Autre. Maximum : 50 tâches, 10 prédécesseurs par tâche, fichier de 2 Mo.')
        st.download_button(
            'Modèle CSV fictif : sept jours',safe_csv(tasks_frame(example_project('sept_jours')['taches'])),
            file_name='modele-fictif-sept-jours.csv',mime='text/csv',icon=':material/download:',
        )
        st.caption('CSV UTF-8 séparé par des virgules. Pour Excel, feuille Taches si elle existe, sinon première feuille. Seuls les fichiers .xlsx sont acceptés.')
    with st.expander('Provenance et état du projet',icon=':material/fact_check:'):
        table_view(pd.DataFrame([
            {'Champ':'Origine déclarée','Valeur':project['origine']},
            {'Champ':'Nombre de tâches','Valeur':str(len(project['taches']))},
            {'Champ':'Simulation active','Valeur':'Oui' if project['simulation_active'] else 'Non'},
            {'Champ':'Comparaison active','Valeur':'Oui' if project['comparaison_active'] else 'Non'},
            {'Champ':'Empreinte SHA-256 du JSON','Valeur':hashlib.sha256(configuration_bytes(project)).hexdigest()},
        ]))
        st.caption('L’empreinte permet de reconnaître une configuration identique, pas d’authentifier son auteur. Aucun historique permanent n’est stocké.')


def configuration_bytes(project):
    return canonical(project).encode('utf-8')


def reference_checks():
    checks = []

    def record(name, expected, obtained, success):
        checks.append({
            'Contrôle':name,
            'Attendu':expected,
            'Obtenu':obtained,
            'État':'Réussi' if success else 'Échec',
        })

    mini = example_project('sept_jours')
    result = compute_schedule(mini)
    record('Durée du mini-lot','7 j',f'{fmt(result["duree_j"])} j',result['duree_j'] == 7.)
    record('Chemin critique','A → B → D',' → '.join(result['chemin']),result['chemin'] == ['A','B','D'])
    margin = float(result['table'].set_index('id').loc['C','marge_j'])
    record('Marge de C','1 j',f'{fmt(margin)} j',margin == 1.)
    simulation = simulate_project(mini)
    record('Dépassement strict à 7 j','0 %',f'{fmt(100*simulation["probabilite"])} %',simulation['probabilite'] == 0.)
    mini['parametres']['echeance_j'] = 6.9
    simulation = simulate_project(mini)
    record('Dépassement à 6,9 j','100 %',f'{fmt(100*simulation["probabilite"])} %',simulation['probabilite'] == 1.)
    chantier = compute_schedule(example_project())
    record('Chantier fictif nominal','18 j',f'{fmt(chantier["duree_j"])} j',chantier['duree_j'] == 18.)
    concurrent = example_project('deux_chemins')
    concurrent['parametres']['tirages'] = 50000
    simulation = simulate_project(concurrent)
    expected = .234375
    record(
        'Deux chemins : P(T > 7)', '23,4375 % ± 1 point',
        f'{fmt(100*simulation["probabilite"],4)} %',
        abs(simulation['probabilite']-expected) < .01,
    )
    invalid = copy.deepcopy(example_project('sept_jours'))
    invalid['taches'][0]['predecesseurs'] = ['D']
    try:
        validate_project(invalid)
        refused = False
    except ValueError:
        refused = True
    record('Boucle de dépendances','Import refusé','Refusé' if refused else 'Accepté',refused)
    return pd.DataFrame(checks)


def glossary():
    return pd.DataFrame([
        {'Mot':'Tâche','Sens simple':'Une action à réaliser, avec une durée : préparer, creuser, assembler…'},
        {'Mot':'Prédécesseur','Sens simple':'Une tâche à terminer avant de pouvoir commencer une autre.'},
        {'Mot':'Lien fin-début','Sens simple':'La suivante attend la fin de la précédente, sans délai supplémentaire.'},
        {'Mot':'Jour abstrait','Sens simple':'Une unité de durée, sans date, week-end ni jour férié.'},
        {'Mot':'Gantt','Sens simple':'Des barres qui montrent quand les tâches commencent et finissent.'},
        {'Mot':'Réseau','Sens simple':'Des boîtes et des flèches qui montrent qui attend qui.'},
        {'Mot':'Chemin critique','Sens simple':'Une suite de tâches qui impose la durée la plus longue du projet.'},
        {'Mot':'Marge totale','Sens simple':'Retard possible d’une tâche sans repousser la fin du projet, toutes choses égales par ailleurs.'},
        {'Mot':'Marge libre','Sens simple':'Retard possible sans repousser le début au plus tôt des tâches suivantes.'},
        {'Mot':'a / m / b','Sens simple':'Durée minimale, durée la plus probable, durée maximale.'},
        {'Mot':'Loi triangulaire','Sens simple':'Un modèle borné défini par ces trois estimations ; pas une loi déduite automatiquement de données.'},
        {'Mot':'Monte-Carlo','Sens simple':'Recalculer le projet de nombreuses fois avec des durées tirées au hasard.'},
        {'Mot':'Graine','Sens simple':'Un nombre qui permet de reproduire une même série de tirages.'},
        {'Mot':'P80','Sens simple':'Le quantile à 80 % des durées simulées, pas une garantie.'},
        {'Mot':'Criticité','Sens simple':'La fréquence à laquelle une tâche est critique dans les plannings simulés.'},
        {'Mot':'Corrélation de rang','Sens simple':'Une mesure d’association entre les grandes ou petites valeurs de deux durées.'},
    ])


def lesson_start():
    st.subheader('Comprendre avec quatre tâches')
    st.write(
        'Imagine un petit travail réalisé par des équipes disponibles. '
        '**A** prépare le travail. Ensuite **B** et **C** peuvent avancer en même temps. '
        '**D** ne commence que quand B et C sont terminées.'
    )
    project = example_project('sept_jours')
    result = compute_schedule(project)
    draw_diagram(network_svg(result),'lesson_network')
    table_view(pd.DataFrame([
        {'Tâche':'A','Durée (j)':2,'Attendre':'Personne','Début (j)':0,'Fin (j)':2},
        {'Tâche':'B','Durée (j)':4,'Attendre':'A','Début (j)':2,'Fin (j)':6},
        {'Tâche':'C','Durée (j)':3,'Attendre':'A','Début (j)':2,'Fin (j)':5},
        {'Tâche':'D','Durée (j)':1,'Attendre':'B et C','Début (j)':6,'Fin (j)':7},
    ]))
    st.markdown(
        '1. A finit au jour **2**.\n'
        '2. B finit au jour **6** et C au jour **5**.\n'
        '3. D attend la plus tardive : **max(6, 5) = 6**.\n'
        '4. D dure un jour : le projet finit au jour **7**.'
    )
    st.success('Durée du projet : 2 + max(4, 3) + 1 = 7 jours. Additionner 2 + 4 + 3 + 1 donnerait 10 jours, ce qui ignore le travail en parallèle.')
    st.write('Le chemin **A → B → D** dure 7 jours. Le chemin **A → C → D** dure 6 jours. C dispose donc ici d’une marge de **1 jour**.')
    st.caption('Exemple intégralement fictif. Les noms des tâches servent à expliquer le calcul, pas à décrire un chantier réel.')


def lesson_probabilities():
    st.subheader('Pourquoi ajouter des probabilités ?')
    st.write(
        'Dire « cette tâche durera 4 jours » masque parfois une incertitude. '
        'On peut plutôt estimer un minimum de 2 jours, un cas le plus probable de '
        '4 jours et un maximum de 6 jours. La loi triangulaire traduit ces trois hypothèses.'
    )
    chart(triangular_figure(task('X','Tâche fictive','Autre',2,4,6)),'lesson_triangle',height=290)
    st.latex(r'\mathbb{E}[D] = \frac{a+m+b}{3}')
    st.write('Pour 2 / 4 / 6 : moyenne = (2 + 4 + 6) / 3 = **4 jours**. La moyenne triangulaire n’est pas la moyenne PERT (a + 4m + b) / 6.')
    st.write(
        'Dans « Deux chemins concurrents », A et D durent chacun 1 jour. '
        'B et C ont toutes deux une estimation 2 / 4 / 6 et sont indépendantes. '
        'Le projet dure **1 + max(B, C) + 1**.'
    )
    st.latex(r'P(T>7)=1-P(B\leq5, C\leq5)=1-F(5)^2')
    st.latex(r'F(5)=1-\frac{(6-5)^2}{(6-2)(6-4)}=\frac{7}{8}')
    st.latex(r'P(T>7)=1-\left(\frac{7}{8}\right)^2=\frac{15}{64}=23{,}4375\%')
    st.write(
        'La simulation doit donner une valeur proche de **23,44 %**, pas exactement '
        'cette valeur à chaque graine. À 50 000 tirages, un écart inférieur à '
        '1 point de pourcentage est un contrôle pédagogique raisonnable pour ce cas.'
    )
    st.warning('Si B et C subissent la même pluie ou utilisent une même équipe, l’indépendance devient discutable. Cette application ne modélise pas ces dépendances.')
    st.write('Même avec des durées moyennes de 4 jours pour B et C, la durée moyenne simulée du projet peut dépasser le planning nominal de 6 jours : le maximum de deux durées n’est pas la moyenne de leurs maximums nominaux.')
    with st.expander('Formules de la loi triangulaire',icon=':material/functions:'):
        st.write('Pour a < b et a < m < b, la densité vaut :')
        st.latex(r'f(d)=\begin{cases}\frac{2(d-a)}{(b-a)(m-a)},&a\leq d\leq m\\\frac{2(b-d)}{(b-a)(b-m)},&m<d\leq b\\0,&\text{sinon}\end{cases}')
        st.write('SciPy traite aussi les modes placés à une borne. Si a = m = b, le programme utilise une durée constante, sans division par zéro.')
        st.latex(r'\mathrm{Var}(D)=\frac{a^2+m^2+b^2-am-ab-mb}{18}')
        st.write('La densité n’est pas une probabilité ponctuelle. Pour une durée continue, la probabilité d’obtenir exactement une valeur donnée est nulle ; une probabilité porte sur un intervalle.')


def lesson_protocols():
    st.subheader('Parcours de test des outils')
    st.markdown(
        '1. Dans le menu latéral, ouvrir **Exemples fictifs**.\n'
        '2. Choisir **Contrôle simple : sept jours**, confirmer le remplacement si demandé et charger.\n'
        '3. Revenir au **Planning** : durée 7 j, chemin A → B → D, marge de C = 1 j.\n'
        '4. Dans **Risque de retard**, garder une échéance de 7 j et simuler : risque 0 %.\n'
        '5. Relancer à 6,9 j : risque 100 %. Le critère de retard est strict.\n'
        '6. Dans **Scénarios**, choisir B et saisir 2 / 2 / 2 : durée nominale de la variante 6 j.\n'
        '7. Dans **Données et rapports**, exporter le JSON et préparer le PDF.\n'
        '8. Créer un projet vide, puis réimporter le JSON : les quatre tâches et les paramètres reviennent.'
    )
    table_view(pd.DataFrame([
        {'Outil':'Gantt','Saisie / action':'Exemple sept jours','Résultat attendu':'A : 0-2 ; B : 2-6 ; C : 2-5 ; D : 6-7'},
        {'Outil':'Dates au plus tard','Saisie / action':'Sélectionner C','Résultat attendu':'Début au plus tard 3 ; fin au plus tard 6 ; marge 1 j'},
        {'Outil':'Parallélisme','Saisie / action':'Ouvrir Travaux en parallèle','Résultat attendu':'Maximum 2 tâches simultanées'},
        {'Outil':'Risque fixe','Saisie / action':'Échéance 7 j puis 6,9 j','Résultat attendu':'0 % puis 100 %'},
        {'Outil':'Quantiles fixes','Saisie / action':'Simuler sept jours','Résultat attendu':'P50 = P80 = P95 = 7 j'},
        {'Outil':'Variante utile','Saisie / action':'B devient 2 / 2 / 2','Résultat attendu':'6 j ; C devient critique'},
        {'Outil':'Variante non critique','Saisie / action':'C devient 1 / 1 / 1','Résultat attendu':'7 j ; pas de réduction de la durée totale'},
        {'Outil':'Risque variable','Saisie / action':'Deux chemins, 50 000 tirages, échéance 7 j','Résultat attendu':'Proche de 23,44 % ; support de 4 à 8 j'},
        {'Outil':'Reproductibilité','Saisie / action':'Mêmes tâches, graine et nombre de tirages','Résultat attendu':'Mêmes durées simulées dans la même version des bibliothèques'},
        {'Outil':'Import JSON','Saisie / action':'Exporter puis restaurer','Résultat attendu':'Configuration identique, y compris paramètres'},
        {'Outil':'CSV / Excel','Saisie / action':'Exporter les tâches et réimporter','Résultat attendu':'Mêmes tâches ; analyses réinitialisées ; origine import déclarée'},
        {'Outil':'PDF','Saisie / action':'Préparer et télécharger','Résultat attendu':'Titre, origine, tâches, planning et analyses actives'},
    ]))
    st.subheader('Tester les erreurs sans perdre ses données')
    table_view(pd.DataFrame([
        {'Essai':'Min. 5, probable 3, max. 6','Comportement attendu':'Refus : les durées ne sont pas ordonnées'},
        {'Essai':'Nom vide','Comportement attendu':'Refus : une tâche doit avoir un nom'},
        {'Essai':'Deux tâches avec le même ID','Comportement attendu':'Refus : identifiant dupliqué'},
        {'Essai':'Prédécesseur Z inexistant','Comportement attendu':'Refus : identifiant inconnu'},
        {'Essai':'A dépend de D dans le mini-lot','Comportement attendu':'Refus : cycle A → B → D → A'},
        {'Essai':'Supprimer A sans retirer les liens','Comportement attendu':'Refus tant que B et C en dépendent'},
        {'Essai':'Échéance 0 avec une durée positive','Comportement attendu':'100 % de retard'},
        {'Essai':'Toutes les durées à 0, échéance 0','Comportement attendu':'0 jour ; risque 0 % ; quantiles nuls'},
    ]))
    st.caption('Les formulaires invalides ne remplacent pas les données déjà enregistrées. Exporter le JSON avant les essais de remplacement ou de suppression.')


def lesson_quiz():
    st.subheader('Vérifier que je sais expliquer')
    revision = st.session_state['tempo_revision']
    questions = [
        (
            'B finit au jour 6 et C au jour 5. Quand peut commencer D si elle attend B et C ?',
            ['Au jour 5','Au jour 6','Au jour 11'],
            'Au jour 6',
            'D doit attendre les deux tâches. On prend le maximum, pas la somme.',
        ),
        (
            'C possède une marge de 1 jour. Si sa durée augmente seule de 0,5 jour, que devient le mini-lot ?',
            ['Il reste à 7 jours','Il passe à 7,5 jours','Il passe à 10 jours'],
            'Il reste à 7 jours',
            'C finit au jour 5,5, toujours avant B au jour 6.',
        ),
        (
            'Une échéance vaut exactement 7 jours et la durée du projet vaut 7 jours. Y a-t-il retard ?',
            ['Oui','Non'],
            'Non',
            'Le retard est défini ici par T > échéance, pas par T ≥ échéance.',
        ),
        (
            'Une simulation donne P80 = 9 jours. Que peut-on dire ?',
            ['Le chantier finira certainement au jour 9','Le quantile à 80 % des durées simulées vaut 9 jours','80 % des tâches durent 9 jours'],
            'Le quantile à 80 % des durées simulées vaut 9 jours',
            'Le P80 porte sur la durée du projet simulé. Il reste conditionnel aux hypothèses.',
        ),
    ]
    with st.form(f'quiz_{revision}',border=False):
        answers = []
        for index,(question,options,_,_) in enumerate(questions):
            answers.append(st.radio(question,options,index=None,key=f'quiz_answer_{revision}_{index}'))
        submitted = st.form_submit_button('Corriger mes réponses',icon=':material/fact_check:',type='primary')
    if submitted:
        score = 0
        for index,(_,_,correct,explanation) in enumerate(questions):
            if answers[index] == correct:
                score += 1
                st.success(f'Question {index+1} : juste. {explanation}')
            else:
                st.info(f'Question {index+1} : {correct}. {explanation}')
        st.write(f'**Résultat : {score} / {len(questions)}**')


def understand_view(project):
    st.header('Comprendre et tester')
    st.write('La question centrale : **quelles tâches fixent la durée du chantier et quel est le risque de finir après l’échéance ?**')
    tabs = st.tabs(['Exemple simple','Probabilités','Tests guidés','Quiz','Méthodes et limites'])
    with tabs[0]:
        lesson_start()
    with tabs[1]:
        lesson_probabilities()
    with tabs[2]:
        lesson_protocols()
        if st.button('Exécuter les contrôles de référence',icon=':material/verified:',type='primary'):
            with st.spinner('Vérification des huit cas de référence…'):
                checks = reference_checks()
            table_view(checks)
            if checks['État'].eq('Réussi').all():
                st.success('Les huit contrôles ont réussi dans cet environnement.')
            else:
                st.error('Au moins un contrôle a échoué : vérifier les bibliothèques et ne pas interpréter les résultats avant correction.')
        st.caption('Ces contrôles ciblés ne remplacent pas une validation complète du logiciel ni l’expertise d’un professionnel.')
    with tabs[3]:
        lesson_quiz()
    with tabs[4]:
        st.subheader('Méthode du chemin critique')
        st.latex(r'ES_i=\max_{p\in Pred(i)}EF_p,\qquad EF_i=ES_i+d_i')
        st.latex(r'LF_i=\min_{s\in Succ(i)}LS_s,\qquad LS_i=LF_i-d_i')
        st.latex(r'MT_i=LS_i-ES_i')
        st.write('Sans prédécesseur, ES = 0. Sans successeur, LF = durée totale du projet. Le réseau doit être orienté et sans cycle. Le calcul est contrôlé par la longueur du plus long chemin pondéré dans NetworkX.')
        st.subheader('Simulation et précision')
        st.write('NumPy génère des nombres uniformes ; SciPy les transforme en durées triangulaires. Chaque tirage refait les passes avant et arrière du réseau. Le chemin critique peut changer entre les tirages.')
        st.write('Le risque est la proportion de tirages où T dépasse strictement l’échéance. Une tolérance de huit unités de dernier rang autour d’un seuil positif évite les faux retards dus aux arrondis flottants. Le seuil zéro reste strictement nul.')
        st.write('L’intervalle à 95 % est binomial exact (Clopper-Pearson). Il n’est pas affiché lorsque la durée totale est constante. Un risque observé à 0 % sur un modèle variable n’est pas une preuve d’impossibilité.')
        st.write('Les quantiles sont calculés par NumPy avec interpolation linéaire. Les moyennes et probabilités reposent sur tous les tirages, même si la courbe cumulée est allégée pour l’affichage.')
        st.subheader('Ce que le modèle ne fait pas')
        st.warning(LIMITS)
        st.write('Il n’équilibre pas les équipes entre tâches. Il ne planifie ni approvisionnements réels, ni séchage réglementaire, ni sécurité, ni coûts. Une durée nulle représente un jalon. Les bornes choisies ne garantissent pas que les durées réelles y resteront.')
        st.write('Les estimations a / m / b sont saisies par l’utilisateur. Aucune base de données locale, norme, exigence réglementaire ou observation réelle n’est présumée.')
        st.subheader('Glossaire')
        table_view(glossary())
        st.subheader('Références')
        st.markdown(
            '- [NetworkX : graphes sans cycle et plus longs chemins](https://networkx.org/documentation/stable/reference/algorithms/dag.html)\n'
            '- [SciPy : loi triangulaire](https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.triang.html)\n'
            '- [SciPy : corrélation de Spearman](https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.spearmanr.html)\n'
            '- [Streamlit : dépendances sur Community Cloud](https://docs.streamlit.io/deploy/streamlit-community-cloud/deploy-your-app/app-dependencies)'
        )


def main():
    st.set_page_config(
        page_title='TempoChantier | Planning et risque',
        page_icon=':material/event_note:',layout='wide',initial_sidebar_state='auto',
        menu_items={'About':'TempoChantier : atelier pédagogique de planning et d’incertitudes.'},
    )
    st.html(app_css())
    if 'tempo_project' not in st.session_state:
        st.session_state['tempo_project'] = new_project()
        st.session_state['tempo_revision'] = 0
        st.session_state['tempo_section'] = 'Planning'
    project = validate_project(st.session_state['tempo_project'])
    sidebar(project)
    brand(project)
    message = st.session_state.pop('tempo_flash',None)
    if message:
        st.success(message,icon=':material/check_circle:')
    views = {
        'Planning':planning_view,
        'Tâches et liens':tasks_view,
        'Risque de retard':risk_view,
        'Scénarios':scenarios_view,
        'Données et rapports':data_view,
        'Comprendre':understand_view,
    }
    section = st.session_state.get('tempo_section','Planning')
    views.get(section,planning_view)(project)
    st.html(
        '<footer class="tempo-footer">TempoChantier · Atelier pédagogique de génie civil'
        '<br>Jours abstraits · Ressources supposées disponibles · Résultats conditionnels aux hypothèses</footer>'
    )


if __name__ == '__main__':
    main()
